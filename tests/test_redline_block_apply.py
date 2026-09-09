#!/usr/bin/env python3
"""
Slice test (TDD) for issue #621: "block-transcript span compiler --
multi-edit paragraphs, footnotes keyed by issue".

## Root problem this proves fixed

Before this slice nothing could compile a PROVEN block transcript
(`scripts/block_transcript.py`, issue #620) into OOXML tracked changes.
The retired quote patcher could only apply document-unique quotes,
one `replace()` at a time, and its footnote pass anchored on the FIRST
`<w:del>` in the paragraph -- so two issues editing one paragraph both
footnoted the first one's deletion (that module's own lines 61-66 document
the bug), a pure insertion was impossible without a spurious `<w:del>`, and
"edit the SECOND of two identical spans" could not be expressed at all.
`scripts/redline_block_apply.py` does not exist before this slice and this
file FAILS on import until it does.

## What this test asserts (mirrors the issue's Acceptance criteria)

  1. Two edits in one paragraph produce two `<w:del>`/`<w:ins>` pairs in
     correct document order, and each issue's `<w:footnoteReference>` sits
     inside THAT ISSUE's own `<w:ins>` -- asserted by REVISION ID (the
     `w:id` the compiler reported for that issue), never by paragraph
     position, which is exactly the assertion the old footnote path would
     fail.
  2. A mid-paragraph PURE insertion splits the run, emits no `<w:del>` at
     all, and the inserted run carries the split run's own `<w:rPr>` (both
     halves of the split keep it too).
  3. Occurrence disambiguation: the same substring appears twice in one
     paragraph and the transcript edits the SECOND; the FIRST is untouched.
  4. Round-trip verification passes on every produced document, and
     `docx-editor` itself can reopen the output and cleanly
     `accept_all()` / `reject_all()` to the expected before/after text --
     not just "the tags are present".
  5. A `delete` span crossing a physical `<w:p>` boundary inside one
     logical block is a structured `spans_physical_paragraph` failure, not
     an attempted cross-paragraph write.
  6. The rationale footnote BODY is itself tracked (issue #615): wrapped in
     a `<w:ins>` carrying `w:id`/`w:author`/`w:date`, so a reject-all leaves
     no orphaned machine-authored commentary behind.
  7. Scope item 1: the namespace-preservation helpers live in
     `scripts/ooxml_util.py` and NOWHERE else -- every writer that rewrites
     `word/document.xml` reaches the same function objects rather than
     carrying a private copy that could drift.
  8. A rejected transcript is a caller-contract violation (`ValueError`):
     an unproven transcript must never reach a writer.

## Fix round 1 (review findings)

  9. A transcript proven against one block never contaminates a LATER
     paragraph that happens to read the same: physical paragraphs are
     resolved by carried-through identity, so a paragraph that fails to
     match its own (stripped) clean text cannot make the resolver slide
     onto an unrelated duplicate and report `applied`.
 10. A paragraph carrying ordinary leading/trailing whitespace still
     applies -- the block's stripped clean text and the live accepted-view
     text are compared on the same derivation, and offsets are translated
     across the leading whitespace, insertion at block offset 0 included.
 11. When one issue owns several edits in one paragraph, its footnote
     reference hangs on its first `<w:ins>` IN DOCUMENT ORDER, matching what
     `inject_issue_footnotes` documents (revision ids are RECORDED in
     descending-offset order, which is not document order).

Uses python-docx (test-only dependency, matching
the repo's other OOXML fixture builders and `tests/redline/test_inplace_patcher_
core.py`) to build small SYNTHETIC fixtures -- never a real document.

Exit codes: 0 = pass, 1 = fail
"""

from __future__ import annotations

import io
import sys
import tempfile
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parent.parent
SCRIPTS_DIR = REPO_ROOT / "scripts"
sys.path.insert(0, str(SCRIPTS_DIR))


def _import_module():
    try:
        import redline_block_apply as _redline_block_apply  # type: ignore

        return _redline_block_apply, None
    except ImportError as exc:
        return None, (
            f"MISSING: scripts/redline_block_apply.py does not exist or fails "
            f"to import ({exc}).\n"
            f"  FIX: implement apply_block_transcript(docx_bytes, proven, *, "
            f"author, timestamp_iso) per issue #621."
        )


redline_block_apply, IMPORT_ERROR = _import_module()

if redline_block_apply is not None:
    import block_transcript  # type: ignore
    import docx_editor  # type: ignore
    import extraction_normalization_stage  # type: ignore
    import ooxml_util  # type: ignore
    import redline_generate  # type: ignore
    import redline_projections  # type: ignore


WORD_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

AUTHOR = "contract-toaster"
TIMESTAMP = "2026-01-01T00:00:00Z"


def _qn(tag: str) -> str:
    return f"{{{WORD_NS}}}{tag}"


# ---------------------------------------------------------------------------
# Synthetic fixture construction (python-docx allowed in tests only)
#
# The heading paragraph carries a real Heading style so
# `clause_boundaries.is_boundary_paragraph_ooxml` starts a logical block at
# it; body paragraphs stay unstyled so they are that block's PHYSICAL
# paragraphs. Nothing here is drawn from any real document.
# ---------------------------------------------------------------------------


def _make_docx(body_texts: list, *, italic: bool = False) -> bytes:
    import docx  # local import: python-docx is a test-only dependency

    document = docx.Document()
    document.add_paragraph("Section 1. Term", style="Heading 1")
    for text in body_texts:
        paragraph = document.add_paragraph()
        run = paragraph.add_run(text)
        if italic:
            # Deliberately NOT bold: `clause_boundaries` treats a short bold
            # single-line paragraph as a heading boundary, which would make
            # this paragraph its own (empty-bodied) block.
            run.italic = True
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def _prove(docx_bytes: bytes, segments: list):
    """Run the real #620 validator over a one-block transcript against the
    fixture's own first block, so every offset the compiler consumes was
    proven, never hand-written."""
    norm = extraction_normalization_stage.extract_and_normalize(docx_bytes)
    assert norm["status"] == "normalized", norm
    block_map = extraction_normalization_stage.build_block_map(norm["paragraphs"])
    block_id = next(iter(block_map))
    proven = block_transcript.validate_block_patches(
        [{"block_id": block_id, "segments": segments}], [], block_map
    )
    return proven, block_id


def _document_root(docx_bytes: bytes) -> ET.Element:
    with zipfile.ZipFile(io.BytesIO(docx_bytes)) as zf:
        return ET.fromstring(zf.read("word/document.xml"))


def _footnotes_root(docx_bytes: bytes):
    with zipfile.ZipFile(io.BytesIO(docx_bytes)) as zf:
        if "word/footnotes.xml" not in zf.namelist():
            return None
        return ET.fromstring(zf.read("word/footnotes.xml"))


def _body_paragraph(docx_bytes: bytes, index: int) -> ET.Element:
    """The `index`-th `<w:p>` in the document (0 = the heading fixture adds)."""
    return list(_document_root(docx_bytes).iter(_qn("p")))[index]


def _revision_children(p: ET.Element) -> list:
    """`(tag, text)` for every direct-child `<w:ins>`/`<w:del>` of `p`, in
    document order -- the order a reviewer reads them in."""
    out = []
    for child in p:
        if child.tag == _qn("del"):
            out.append(("del", "".join(t.text or "" for t in child.iter(_qn("delText")))))
        elif child.tag == _qn("ins"):
            out.append(("ins", "".join(t.text or "" for t in child.iter(_qn("t")))))
    return out


def _reopen(docx_bytes: bytes, tmp_dir: Path, name: str):
    work = tmp_dir / name
    work.mkdir()
    path = work / "doc.docx"
    path.write_bytes(docx_bytes)
    doc = docx_editor.Document.open(path, author=AUTHOR, workspace_dir=str(work / "ws"))
    return doc, path


# ---------------------------------------------------------------------------
# AC 1 + AC 4: two edits in one paragraph, footnotes anchored by revision id
# ---------------------------------------------------------------------------

_TWO_EDIT_TEXT = (
    "The Term shall be sixty (60) days and the Notice Period shall be ten (10) days."
)
_TWO_EDIT_SEGMENTS = [
    {"op": "keep", "text": "The Term shall be "},
    {"op": "delete", "text": "sixty (60)", "issue_key": "TERM-1"},
    {"op": "insert", "text": "thirty (30)", "issue_key": "TERM-1"},
    {"op": "keep", "text": " days and the Notice Period shall be "},
    {"op": "delete", "text": "ten (10)", "issue_key": "NOTICE-1"},
    {"op": "insert", "text": "thirty (30)", "issue_key": "NOTICE-1"},
    {"op": "keep", "text": " days."},
]
_RATIONALES = {
    "TERM-1": "Term shortened to the negotiated ceiling.",
    "NOTICE-1": "Notice period aligned with the term.",
}


def test_two_edits_in_one_paragraph(failures: list) -> None:
    case = "two_edits_in_one_paragraph"
    docx_bytes = _make_docx([_TWO_EDIT_TEXT])
    proven, _ = _prove(docx_bytes, _TWO_EDIT_SEGMENTS)
    if proven["status"] != "proven":
        failures.append(f"[{case}] fixture transcript did not prove: {proven['failures']!r}")
        return

    result = redline_block_apply.apply_block_transcript(
        docx_bytes,
        proven,
        author=AUTHOR,
        timestamp_iso=TIMESTAMP,
        rationale_by_issue=_RATIONALES,
    )
    if result["failures"]:
        failures.append(f"[{case}] unexpected failures: {result['failures']!r}")
    out = result["docx_bytes"]
    if not out:
        failures.append(f"[{case}] no docx_bytes returned")
        return

    # AC 4: round-trip verification on every produced document.
    try:
        redline_generate.verify_docx_round_trip(out)
    except ValueError as exc:
        failures.append(f"[{case}] verify_docx_round_trip raised: {exc}")

    # AC 1a: two <w:del>/<w:ins> pairs, in the order they appear in the text.
    paragraph = _body_paragraph(out, 1)
    revisions = _revision_children(paragraph)
    expected = [
        ("del", "sixty (60)"),
        ("ins", "thirty (30)"),
        ("del", "ten (10)"),
        ("ins", "thirty (30)"),
    ]
    if revisions != expected:
        failures.append(f"[{case}] revision sequence {revisions!r} != {expected!r}")

    # AC 1b: each issue's footnote reference sits inside THAT issue's own
    # <w:ins>, matched by revision id -- the assertion the old
    # `p.find(w:del)` footnote path cannot satisfy.
    by_issue = result["revision_ids_by_issue"]
    for issue_key in ("TERM-1", "NOTICE-1"):
        if issue_key not in by_issue or not by_issue[issue_key]:
            failures.append(f"[{case}] no revision ids reported for {issue_key}")
            return

    ins_by_id = {}
    for ins in paragraph.iter(_qn("ins")):
        raw = ins.get(_qn("id"))
        if raw is not None:
            ins_by_id[int(raw)] = ins

    footnotes_root = _footnotes_root(out)
    if footnotes_root is None:
        failures.append(f"[{case}] no word/footnotes.xml in the output")
        return
    footnote_text_by_id = {}
    for fn in footnotes_root.findall(_qn("footnote")):
        footnote_text_by_id[fn.get(_qn("id"))] = "".join(
            t.text or "" for t in fn.iter(_qn("t"))
        )

    for issue_key, expected_insert in (("TERM-1", "thirty (30)"), ("NOTICE-1", "thirty (30)")):
        owned_ins = [ins_by_id[rid] for rid in by_issue[issue_key] if rid in ins_by_id]
        if len(owned_ins) != 1:
            failures.append(
                f"[{case}] expected exactly one <w:ins> for {issue_key}, got {len(owned_ins)}"
            )
            continue
        ins = owned_ins[0]
        ins_text = "".join(t.text or "" for t in ins.iter(_qn("t")))
        if ins_text != expected_insert:
            failures.append(f"[{case}] {issue_key} <w:ins> text {ins_text!r}")
        refs = [r.get(_qn("id")) for r in ins.iter(_qn("footnoteReference"))]
        if len(refs) != 1:
            failures.append(
                f"[{case}] {issue_key}: expected 1 footnoteReference inside its own "
                f"<w:ins>, got {refs!r}"
            )
            continue
        body = footnote_text_by_id.get(refs[0], "")
        if _RATIONALES[issue_key] not in body:
            failures.append(
                f"[{case}] {issue_key}'s footnote body is {body!r} -- MISATTRIBUTED "
                f"(expected {_RATIONALES[issue_key]!r})"
            )

    # AC 4: docx-editor reopens the output and accepts/rejects cleanly.
    with tempfile.TemporaryDirectory() as tmp:
        tmp_path = Path(tmp)
        doc, path = _reopen(out, tmp_path, "accept")
        try:
            doc.accept_all()
            doc.save()
        finally:
            doc.close()
        accepted = "".join(
            t.text or "" for t in _body_paragraph(path.read_bytes(), 1).iter(_qn("t"))
        )
        expected_accepted = (
            "The Term shall be thirty (30) days and the Notice Period shall be "
            "thirty (30) days."
        )
        if accepted != expected_accepted:
            failures.append(f"[{case}] accept_all text {accepted!r}")

        doc, path = _reopen(out, tmp_path, "reject")
        try:
            doc.reject_all()
            doc.save()
        finally:
            doc.close()
        rejected = "".join(
            t.text or "" for t in _body_paragraph(path.read_bytes(), 1).iter(_qn("t"))
        )
        if rejected != _TWO_EDIT_TEXT:
            failures.append(f"[{case}] reject_all text {rejected!r} != {_TWO_EDIT_TEXT!r}")


# ---------------------------------------------------------------------------
# AC 2: a mid-paragraph pure insertion
# ---------------------------------------------------------------------------

_INSERT_TEXT = "The Receiving Party shall protect the Confidential Information."


def test_pure_insertion_splits_run_and_keeps_rpr(failures: list) -> None:
    case = "pure_insertion_splits_run_and_keeps_rpr"
    docx_bytes = _make_docx([_INSERT_TEXT], italic=True)
    proven, _ = _prove(
        docx_bytes,
        [
            {"op": "keep", "text": "The Receiving Party shall "},
            {"op": "insert", "text": "use reasonable care to ", "issue_key": "CONF-1"},
            {"op": "keep", "text": "protect the Confidential Information."},
        ],
    )
    if proven["status"] != "proven":
        failures.append(f"[{case}] fixture transcript did not prove: {proven['failures']!r}")
        return

    result = redline_block_apply.apply_block_transcript(
        docx_bytes, proven, author=AUTHOR, timestamp_iso=TIMESTAMP
    )
    if result["failures"]:
        failures.append(f"[{case}] unexpected failures: {result['failures']!r}")
    out = result["docx_bytes"]
    if not out:
        failures.append(f"[{case}] no docx_bytes returned")
        return
    try:
        redline_generate.verify_docx_round_trip(out)
    except ValueError as exc:
        failures.append(f"[{case}] verify_docx_round_trip raised: {exc}")

    paragraph = _body_paragraph(out, 1)

    # No spurious deletion: this is the whole point of a pure insertion.
    dels = list(paragraph.iter(_qn("del")))
    if dels:
        failures.append(f"[{case}] a pure insertion emitted {len(dels)} <w:del> element(s)")

    revisions = _revision_children(paragraph)
    if revisions != [("ins", "use reasonable care to ")]:
        failures.append(f"[{case}] revision sequence {revisions!r}")
        return

    # The run was SPLIT: the paragraph now carries the text before the
    # insertion and the text after it as separate plain runs.
    plain_runs = [
        "".join(t.text or "" for t in child.iter(_qn("t")))
        for child in paragraph
        if child.tag == _qn("r")
    ]
    if plain_runs != ["The Receiving Party shall ", "protect the Confidential Information."]:
        failures.append(f"[{case}] the run was not split as expected: {plain_runs!r}")

    # Every run involved -- both halves of the split AND the inserted run --
    # carries the split run's own <w:rPr> (italic), so formatting continues
    # across the insertion.
    def _has_italic(run: ET.Element) -> bool:
        rpr = run.find(_qn("rPr"))
        return rpr is not None and rpr.find(_qn("i")) is not None

    for child in paragraph:
        if child.tag == _qn("r") and not _has_italic(child):
            text = "".join(t.text or "" for t in child.iter(_qn("t")))
            failures.append(f"[{case}] split half {text!r} lost its <w:rPr>")
    ins_el = next(paragraph.iter(_qn("ins")))
    inserted_runs = [r for r in ins_el if r.tag == _qn("r")]
    if not inserted_runs or not _has_italic(inserted_runs[0]):
        failures.append(f"[{case}] the inserted run does not carry the split run's <w:rPr>")

    if ins_el.get(_qn("author")) != AUTHOR or ins_el.get(_qn("date")) != TIMESTAMP:
        failures.append(
            f"[{case}] <w:ins> author/date "
            f"{ins_el.get(_qn('author'))!r}/{ins_el.get(_qn('date'))!r}"
        )

    with tempfile.TemporaryDirectory() as tmp:
        doc, path = _reopen(out, Path(tmp), "reject")
        try:
            doc.reject_all()
            doc.save()
        finally:
            doc.close()
        rejected = "".join(
            t.text or "" for t in _body_paragraph(path.read_bytes(), 1).iter(_qn("t"))
        )
        if rejected != _INSERT_TEXT:
            failures.append(f"[{case}] reject_all text {rejected!r} != {_INSERT_TEXT!r}")


# ---------------------------------------------------------------------------
# AC 3: occurrence disambiguation
# ---------------------------------------------------------------------------

_REPEATED_TEXT = (
    "The Term shall be thirty (30) days. The Notice Period shall be thirty (30) days."
)


def test_occurrence_disambiguation_edits_the_second_match(failures: list) -> None:
    case = "occurrence_disambiguation_edits_the_second_match"
    docx_bytes = _make_docx([_REPEATED_TEXT])
    proven, _ = _prove(
        docx_bytes,
        [
            {
                "op": "keep",
                "text": "The Term shall be thirty (30) days. The Notice Period shall be ",
            },
            {"op": "delete", "text": "thirty (30)", "issue_key": "NOTICE-1"},
            {"op": "insert", "text": "sixty (60)", "issue_key": "NOTICE-1"},
            {"op": "keep", "text": " days."},
        ],
    )
    if proven["status"] != "proven":
        failures.append(f"[{case}] fixture transcript did not prove: {proven['failures']!r}")
        return

    result = redline_block_apply.apply_block_transcript(
        docx_bytes, proven, author=AUTHOR, timestamp_iso=TIMESTAMP
    )
    out = result["docx_bytes"]
    if result["failures"] or not out:
        failures.append(f"[{case}] failures={result['failures']!r} bytes={bool(out)}")
        return
    try:
        redline_generate.verify_docx_round_trip(out)
    except ValueError as exc:
        failures.append(f"[{case}] verify_docx_round_trip raised: {exc}")

    paragraph = _body_paragraph(out, 1)
    revisions = _revision_children(paragraph)
    if revisions != [("del", "thirty (30)"), ("ins", "sixty (60)")]:
        failures.append(f"[{case}] revision sequence {revisions!r}")

    # The FIRST occurrence is untouched: it is still plain, un-revised text
    # sitting BEFORE the deletion.
    leading_plain = "".join(
        "".join(t.text or "" for t in child.iter(_qn("t")))
        for child in paragraph
        if child.tag == _qn("r")
    )
    if "The Term shall be thirty (30) days." not in leading_plain:
        failures.append(
            f"[{case}] the FIRST occurrence was not left untouched: plain text is "
            f"{leading_plain!r}"
        )

    with tempfile.TemporaryDirectory() as tmp:
        doc, path = _reopen(out, Path(tmp), "accept")
        try:
            doc.accept_all()
            doc.save()
        finally:
            doc.close()
        accepted = "".join(
            t.text or "" for t in _body_paragraph(path.read_bytes(), 1).iter(_qn("t"))
        )
        expected = (
            "The Term shall be thirty (30) days. The Notice Period shall be sixty (60) days."
        )
        if accepted != expected:
            failures.append(f"[{case}] accept_all text {accepted!r} != {expected!r}")


# ---------------------------------------------------------------------------
# Descending application order: BOTH occurrences of one substring edited
#
# This is what makes "apply in descending offset order" load-bearing rather
# than stylistic. `docx_editor` edits are addressed by (text, occurrence)
# against the paragraph's CURRENT accepted view, and every occurrence index
# here was computed against the ORIGINAL text. Applying the lower offset
# first removes an occurrence from that view, so the higher edit's index
# then addresses the wrong span -- or nothing at all. Descending order never
# disturbs the text before the edit, so every remaining index stays true.
# ---------------------------------------------------------------------------


def test_both_occurrences_edited_needs_descending_order(failures: list) -> None:
    case = "both_occurrences_edited_needs_descending_order"
    docx_bytes = _make_docx([_REPEATED_TEXT])
    proven, _ = _prove(
        docx_bytes,
        [
            {"op": "keep", "text": "The Term shall be "},
            {"op": "delete", "text": "thirty (30)", "issue_key": "TERM-1"},
            {"op": "insert", "text": "sixty (60)", "issue_key": "TERM-1"},
            {"op": "keep", "text": " days. The Notice Period shall be "},
            {"op": "delete", "text": "thirty (30)", "issue_key": "NOTICE-1"},
            {"op": "insert", "text": "ninety (90)", "issue_key": "NOTICE-1"},
            {"op": "keep", "text": " days."},
        ],
    )
    if proven["status"] != "proven":
        failures.append(f"[{case}] fixture transcript did not prove: {proven['failures']!r}")
        return

    result = redline_block_apply.apply_block_transcript(
        docx_bytes, proven, author=AUTHOR, timestamp_iso=TIMESTAMP
    )
    if result["failures"]:
        failures.append(f"[{case}] unexpected failures: {result['failures']!r}")
    out = result["docx_bytes"]
    if not out:
        failures.append(f"[{case}] no docx_bytes returned")
        return
    try:
        redline_generate.verify_docx_round_trip(out)
    except ValueError as exc:
        failures.append(f"[{case}] verify_docx_round_trip raised: {exc}")

    revisions = _revision_children(_body_paragraph(out, 1))
    expected = [
        ("del", "thirty (30)"),
        ("ins", "sixty (60)"),
        ("del", "thirty (30)"),
        ("ins", "ninety (90)"),
    ]
    if revisions != expected:
        failures.append(f"[{case}] revision sequence {revisions!r} != {expected!r}")

    with tempfile.TemporaryDirectory() as tmp:
        doc, path = _reopen(out, Path(tmp), "accept")
        try:
            doc.accept_all()
            doc.save()
        finally:
            doc.close()
        accepted = "".join(
            t.text or "" for t in _body_paragraph(path.read_bytes(), 1).iter(_qn("t"))
        )
        want = (
            "The Term shall be sixty (60) days. The Notice Period shall be ninety (90) days."
        )
        if accepted != want:
            failures.append(f"[{case}] accept_all text {accepted!r} != {want!r}")


# ---------------------------------------------------------------------------
# A replacement AND a pure insertion in one paragraph
#
# The two are written by different passes (docx_editor, then owned XML), so
# the insertion's proven offset has to be translated across what the
# replacement already changed in the accepted view. Getting that translation
# wrong lands the insertion inside the neighbouring word.
# ---------------------------------------------------------------------------

_MIXED_TEXT = "The Term shall be sixty (60) days and the Agreement shall renew automatically."


def test_replacement_and_pure_insertion_in_one_paragraph(failures: list) -> None:
    case = "replacement_and_pure_insertion_in_one_paragraph"
    docx_bytes = _make_docx([_MIXED_TEXT])
    proven, _ = _prove(
        docx_bytes,
        [
            {"op": "keep", "text": "The Term shall be "},
            {"op": "delete", "text": "sixty (60)", "issue_key": "TERM-1"},
            {"op": "insert", "text": "thirty (30)", "issue_key": "TERM-1"},
            {"op": "keep", "text": " days and the Agreement shall "},
            {"op": "insert", "text": "not ", "issue_key": "RENEW-1"},
            {"op": "keep", "text": "renew automatically."},
        ],
    )
    if proven["status"] != "proven":
        failures.append(f"[{case}] fixture transcript did not prove: {proven['failures']!r}")
        return

    result = redline_block_apply.apply_block_transcript(
        docx_bytes,
        proven,
        author=AUTHOR,
        timestamp_iso=TIMESTAMP,
        rationale_by_issue={
            "TERM-1": "Term shortened to the negotiated ceiling.",
            "RENEW-1": "Automatic renewal removed.",
        },
    )
    if result["failures"]:
        failures.append(f"[{case}] unexpected failures: {result['failures']!r}")
    out = result["docx_bytes"]
    if not out:
        failures.append(f"[{case}] no docx_bytes returned")
        return
    try:
        redline_generate.verify_docx_round_trip(out)
    except ValueError as exc:
        failures.append(f"[{case}] verify_docx_round_trip raised: {exc}")

    paragraph = _body_paragraph(out, 1)
    revisions = _revision_children(paragraph)
    expected = [("del", "sixty (60)"), ("ins", "thirty (30)"), ("ins", "not ")]
    if revisions != expected:
        failures.append(f"[{case}] revision sequence {revisions!r} != {expected!r}")

    # Exactly one <w:del>: the insertion did NOT drag a spurious deletion
    # along with it, which is what the quote path had to do.
    if len(list(paragraph.iter(_qn("del")))) != 1:
        failures.append(f"[{case}] expected exactly one <w:del> in the paragraph")

    with tempfile.TemporaryDirectory() as tmp:
        tmp_path = Path(tmp)
        doc, path = _reopen(out, tmp_path, "accept")
        try:
            doc.accept_all()
            doc.save()
        finally:
            doc.close()
        accepted = "".join(
            t.text or "" for t in _body_paragraph(path.read_bytes(), 1).iter(_qn("t"))
        )
        want = (
            "The Term shall be thirty (30) days and the Agreement shall not renew "
            "automatically."
        )
        if accepted != want:
            failures.append(f"[{case}] accept_all text {accepted!r} != {want!r}")

        doc, path = _reopen(out, tmp_path, "reject")
        try:
            doc.reject_all()
            doc.save()
        finally:
            doc.close()
        rejected = "".join(
            t.text or "" for t in _body_paragraph(path.read_bytes(), 1).iter(_qn("t"))
        )
        if rejected != _MIXED_TEXT:
            failures.append(f"[{case}] reject_all text {rejected!r} != {_MIXED_TEXT!r}")


# ---------------------------------------------------------------------------
# Insertions on the paragraph's outer boundaries
#
# Offset 0 has no run to its left and the end offset has none to its right,
# so both take the boundary branch rather than the split branch -- and both
# are shapes `docx_editor`'s anchor-text `insert_before`/`insert_after`
# cannot address at all.
# ---------------------------------------------------------------------------

_BOUNDARY_TEXT = "The Term is thirty (30) days."


def test_insertions_at_both_paragraph_boundaries(failures: list) -> None:
    case = "insertions_at_both_paragraph_boundaries"
    docx_bytes = _make_docx([_BOUNDARY_TEXT], italic=True)
    proven, _ = _prove(
        docx_bytes,
        [
            {"op": "insert", "text": "Subject to Section 9, ", "issue_key": "HEAD-1"},
            {"op": "keep", "text": _BOUNDARY_TEXT},
            {"op": "insert", "text": " No extension applies.", "issue_key": "TAIL-1"},
        ],
    )
    if proven["status"] != "proven":
        failures.append(f"[{case}] fixture transcript did not prove: {proven['failures']!r}")
        return

    result = redline_block_apply.apply_block_transcript(
        docx_bytes, proven, author=AUTHOR, timestamp_iso=TIMESTAMP
    )
    if result["failures"]:
        failures.append(f"[{case}] unexpected failures: {result['failures']!r}")
    out = result["docx_bytes"]
    if not out:
        failures.append(f"[{case}] no docx_bytes returned")
        return
    try:
        redline_generate.verify_docx_round_trip(out)
    except ValueError as exc:
        failures.append(f"[{case}] verify_docx_round_trip raised: {exc}")

    paragraph = _body_paragraph(out, 1)
    revisions = _revision_children(paragraph)
    expected = [("ins", "Subject to Section 9, "), ("ins", " No extension applies.")]
    if revisions != expected:
        failures.append(f"[{case}] revision sequence {revisions!r} != {expected!r}")
    if list(paragraph.iter(_qn("del"))):
        failures.append(f"[{case}] a pure insertion emitted a <w:del>")

    with tempfile.TemporaryDirectory() as tmp:
        tmp_path = Path(tmp)
        doc, path = _reopen(out, tmp_path, "accept")
        try:
            doc.accept_all()
            doc.save()
        finally:
            doc.close()
        accepted = "".join(
            t.text or "" for t in _body_paragraph(path.read_bytes(), 1).iter(_qn("t"))
        )
        want = "Subject to Section 9, The Term is thirty (30) days. No extension applies."
        if accepted != want:
            failures.append(f"[{case}] accept_all text {accepted!r} != {want!r}")

        doc, path = _reopen(out, tmp_path, "reject")
        try:
            doc.reject_all()
            doc.save()
        finally:
            doc.close()
        rejected = "".join(
            t.text or "" for t in _body_paragraph(path.read_bytes(), 1).iter(_qn("t"))
        )
        if rejected != _BOUNDARY_TEXT:
            failures.append(f"[{case}] reject_all text {rejected!r} != {_BOUNDARY_TEXT!r}")


# ---------------------------------------------------------------------------
# Scope item 2: a delete crossing a physical <w:p> boundary fails closed
# ---------------------------------------------------------------------------


def test_delete_across_physical_paragraphs_fails_closed(failures: list) -> None:
    case = "delete_across_physical_paragraphs_fails_closed"
    docx_bytes = _make_docx(
        ["First physical paragraph.", "Second physical paragraph."]
    )
    proven, block_id = _prove(
        docx_bytes,
        [
            {"op": "keep", "text": "First physical "},
            {"op": "delete", "text": "paragraph.\nSecond physical", "issue_key": "X-1"},
            {"op": "keep", "text": " paragraph."},
        ],
    )
    if proven["status"] != "proven":
        failures.append(f"[{case}] fixture transcript did not prove: {proven['failures']!r}")
        return

    result = redline_block_apply.apply_block_transcript(
        docx_bytes, proven, author=AUTHOR, timestamp_iso=TIMESTAMP
    )
    if result["docx_bytes"] is not None:
        failures.append(f"[{case}] a cross-paragraph delete produced bytes anyway")
    if result["applied"]:
        failures.append(f"[{case}] a cross-paragraph delete was applied: {result['applied']!r}")
    reasons = [f["reason"] for f in result["failures"]]
    if reasons != [redline_block_apply.REASON_SPANS_PHYSICAL_PARAGRAPH]:
        failures.append(f"[{case}] failure reasons {reasons!r}")
        return
    if result["failures"][0].get("block_id") != block_id:
        failures.append(f"[{case}] failure is not keyed to the block: {result['failures'][0]!r}")


# ---------------------------------------------------------------------------
# Issue #615: the footnote BODY is tracked too
# ---------------------------------------------------------------------------


def test_footnote_body_is_itself_tracked(failures: list) -> None:
    case = "footnote_body_is_itself_tracked"
    docx_bytes = _make_docx([_TWO_EDIT_TEXT])
    proven, _ = _prove(docx_bytes, _TWO_EDIT_SEGMENTS)
    result = redline_block_apply.apply_block_transcript(
        docx_bytes,
        proven,
        author=AUTHOR,
        timestamp_iso=TIMESTAMP,
        rationale_by_issue=_RATIONALES,
    )
    out = result["docx_bytes"]
    if not out:
        failures.append(f"[{case}] no docx_bytes returned")
        return

    footnotes_root = _footnotes_root(out)
    if footnotes_root is None:
        failures.append(f"[{case}] no word/footnotes.xml in the output")
        return

    for fn in footnotes_root.findall(_qn("footnote")):
        body = "".join(t.text or "" for t in fn.iter(_qn("t")))
        if not body.strip():
            continue  # the two mandatory separator footnotes carry no text
        tracked = [ins for ins in fn.iter(_qn("ins"))]
        if not tracked:
            failures.append(
                f"[{case}] footnote {fn.get(_qn('id'))} body {body!r} is UNTRACKED -- "
                f"reject-all would leave it orphaned (issue #615)"
            )
            continue
        ins = tracked[0]
        for attr in ("id", "author", "date"):
            if ins.get(_qn(attr)) in (None, ""):
                failures.append(f"[{case}] footnote body <w:ins> is missing w:{attr}")
        if ins.get(_qn("author")) != AUTHOR or ins.get(_qn("date")) != TIMESTAMP:
            failures.append(
                f"[{case}] footnote body <w:ins> author/date "
                f"{ins.get(_qn('author'))!r}/{ins.get(_qn('date'))!r}"
            )
        # The text really is INSIDE the tracked wrapper, not merely beside it.
        if body.strip() not in "".join(t.text or "" for t in ins.iter(_qn("t"))):
            failures.append(f"[{case}] footnote text sits outside the <w:ins> wrapper")

    # Programmatic reject-all over word/footnotes.xml: dropping every
    # <w:ins> subtree must leave no rationale text behind.
    surviving = _strip_tracked_insertions(footnotes_root)
    for rationale in _RATIONALES.values():
        if rationale in surviving:
            failures.append(
                f"[{case}] rejecting all revisions still leaves {rationale!r} in the "
                f"footnotes part"
            )


def _strip_tracked_insertions(root: ET.Element) -> str:
    """Simulate Word's reject-all over one part: remove every `<w:ins>`
    subtree and return whatever visible text survives."""

    def prune(node: ET.Element) -> None:
        for child in list(node):
            if child.tag == _qn("ins"):
                node.remove(child)
            else:
                prune(child)

    prune(root)
    return "".join(t.text or "" for t in root.iter(_qn("t")))


# ---------------------------------------------------------------------------
# Scope item 1: the shared ooxml_util extraction is behaviour-neutral
# ---------------------------------------------------------------------------


def test_ooxml_util_is_the_single_definition(failures: list) -> None:
    case = "ooxml_util_is_the_single_definition"
    shared_names = (
        "scan_tag_end",
        "root_open_tag",
        "declared_namespaces",
        "register_declared_namespaces",
        "declared_namespaces_anywhere",
        "merge_hoisted_namespaces",
    )
    # Every module that rewrites word/document.xml. If one of them ever grows
    # its own same-named helper, this catches the copy before it can drift.
    writers = (
        redline_block_apply,
        redline_generate,
        redline_projections,
        extraction_normalization_stage,
    )
    for shared_name in shared_names:
        shared = getattr(ooxml_util, shared_name, None)
        if shared is None:
            failures.append(f"[{case}] ooxml_util.{shared_name} does not exist")
            continue
        if getattr(shared, "__module__", None) != "ooxml_util":
            failures.append(
                f"[{case}] ooxml_util.{shared_name} is defined in "
                f"{getattr(shared, '__module__', None)!r}, not in ooxml_util"
            )
        for writer in writers:
            local = getattr(writer, shared_name, None)
            if local is not None and local is not shared:
                failures.append(
                    f"[{case}] {writer.__name__}.{shared_name} is not "
                    f"ooxml_util.{shared_name} -- the helper was copied, not shared"
                )

    # And it still does its job: an xmlns declared only inside an attribute
    # VALUE survives the splice, which is the property the helpers exist for.
    open_tag = (
        '<w:document xmlns:w="%s" xmlns:mc="urn:example:mc" mc:Ignorable="w14 wp14">' % WORD_NS
    )
    auto = '<w:document xmlns:w="%s" xmlns:a="urn:example:hoisted">' % WORD_NS
    merged = ooxml_util.merge_hoisted_namespaces(open_tag, auto)
    if 'mc:Ignorable="w14 wp14"' not in merged or "urn:example:hoisted" not in merged:
        failures.append(f"[{case}] merge_hoisted_namespaces lost a declaration: {merged!r}")


# ---------------------------------------------------------------------------
# Caller contract: an unproven transcript never reaches a writer
# ---------------------------------------------------------------------------


def test_rejected_transcript_raises(failures: list) -> None:
    case = "rejected_transcript_raises"
    docx_bytes = _make_docx([_TWO_EDIT_TEXT])
    proven, _ = _prove(
        docx_bytes,
        [
            # A silently rewritten `keep` -- #620's unmarked-rewrite guard.
            {"op": "keep", "text": "The Term shall be, "},
            {"op": "delete", "text": "sixty (60)", "issue_key": "TERM-1"},
            {"op": "keep", "text": " days and the Notice Period shall be ten (10) days."},
        ],
    )
    if proven["status"] != "rejected":
        failures.append(f"[{case}] fixture transcript unexpectedly proved: {proven!r}")
        return
    try:
        redline_block_apply.apply_block_transcript(
            docx_bytes, proven, author=AUTHOR, timestamp_iso=TIMESTAMP
        )
    except ValueError:
        return
    failures.append(f"[{case}] a rejected transcript did NOT raise -- it reached the writer")


# ---------------------------------------------------------------------------
# Fix round 1, finding 1: identity resolution -- an edit proven against ONE
# block must never slide onto a later paragraph that happens to read the same
# ---------------------------------------------------------------------------


def _make_two_section_docx(pairs: list) -> bytes:
    """`[(heading, body), ...]` -- each heading starts its own logical block,
    each body is that block's single physical paragraph. Synthetic only."""
    import docx  # local import: python-docx is a test-only dependency

    document = docx.Document()
    for heading, body in pairs:
        document.add_paragraph(heading, style="Heading 1")
        paragraph = document.add_paragraph()
        paragraph.add_run(body)
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


# The two bodies differ ONLY by a trailing space -- ordinary in real Word
# documents, and enough to make the first paragraph fail to match its own
# (stripped) clean text. A forward text scan then does not stop: it slides
# onto Section 2's identical paragraph and writes the tracked change into
# the WRONG CLAUSE while reporting `applied` with zero failures.
_DUPLICATE_BODY = "Either party may terminate on ten (10) days notice."
_DUPLICATE_SECTIONS = [
    ("Section 1. Term", _DUPLICATE_BODY + " "),
    ("Section 2. Termination", _DUPLICATE_BODY),
]


def test_edit_never_slides_onto_a_later_equal_paragraph(failures: list) -> None:
    case = "edit_never_slides_onto_a_later_equal_paragraph"
    docx_bytes = _make_two_section_docx(_DUPLICATE_SECTIONS)
    proven, block_id = _prove(
        docx_bytes,
        [
            {"op": "keep", "text": "Either party may terminate on "},
            {"op": "delete", "text": "ten (10)", "issue_key": "TERM-1"},
            {"op": "insert", "text": "thirty (30)", "issue_key": "TERM-1"},
            {"op": "keep", "text": " days notice."},
        ],
    )
    if proven["status"] != "proven":
        failures.append(f"[{case}] fixture transcript did not prove: {proven['failures']!r}")
        return
    if block_id != "p0001":
        failures.append(f"[{case}] fixture proved against {block_id!r}, expected the first block")
        return

    result = redline_block_apply.apply_block_transcript(
        docx_bytes, proven, author=AUTHOR, timestamp_iso=TIMESTAMP
    )
    out = result["docx_bytes"]
    if out is None:
        # Fail-closed is acceptable here; writing into Section 2 is not.
        reasons = [f["reason"] for f in result["failures"]]
        if reasons != [redline_block_apply.REASON_PARAGRAPH_NOT_RESOLVED]:
            failures.append(f"[{case}] no bytes and unexpected reasons {reasons!r}")
        return

    # THE assertion: Section 2's paragraph is untouched. `<w:p>` 3 is
    # Section 2's body (0 = "Section 1. Term", 1 = its body, 2 = "Section 2.
    # Termination", 3 = its body).
    contaminated = _revision_children(_body_paragraph(out, 3))
    if contaminated:
        failures.append(
            f"[{case}] the edit landed in SECTION 2's paragraph: {contaminated!r} -- "
            f"a transcript proven against {block_id!r} contaminated another clause"
        )

    # ... and it landed in Section 1's own paragraph instead.
    edited = _revision_children(_body_paragraph(out, 1))
    expected = [("del", "ten (10)"), ("ins", "thirty (30)")]
    if edited != expected:
        failures.append(f"[{case}] Section 1 revision sequence {edited!r} != {expected!r}")
    if result["failures"]:
        failures.append(f"[{case}] unexpected failures: {result['failures']!r}")
    try:
        redline_generate.verify_docx_round_trip(out)
    except ValueError as exc:
        failures.append(f"[{case}] verify_docx_round_trip raised: {exc}")


# ---------------------------------------------------------------------------
# Fix round 1, finding 2: a paragraph carrying leading/trailing whitespace
# ---------------------------------------------------------------------------

_WHITESPACE_BODY = "  The Term shall be sixty (60) days.  "


def test_whitespace_padded_paragraph_still_applies(failures: list) -> None:
    case = "whitespace_padded_paragraph_still_applies"
    docx_bytes = _make_docx([_WHITESPACE_BODY])
    proven, _ = _prove(
        docx_bytes,
        [
            # A pure insertion at block offset 0 -- the strongest test that
            # offsets are translated across the paragraph's leading
            # whitespace rather than applied to the stripped text.
            {"op": "insert", "text": "Subject to Section 9, ", "issue_key": "TERM-1"},
            {"op": "keep", "text": "The Term shall be "},
            {"op": "delete", "text": "sixty (60)", "issue_key": "TERM-1"},
            {"op": "insert", "text": "thirty (30)", "issue_key": "TERM-1"},
            {"op": "keep", "text": " days."},
        ],
    )
    if proven["status"] != "proven":
        failures.append(f"[{case}] fixture transcript did not prove: {proven['failures']!r}")
        return

    result = redline_block_apply.apply_block_transcript(
        docx_bytes, proven, author=AUTHOR, timestamp_iso=TIMESTAMP
    )
    if result["failures"]:
        failures.append(
            f"[{case}] a paragraph with ordinary leading/trailing whitespace failed: "
            f"{result['failures']!r}"
        )
    out = result["docx_bytes"]
    if not out:
        failures.append(f"[{case}] no docx_bytes returned -- every edit in the block was lost")
        return
    if len(result["applied"]) != 2:
        failures.append(f"[{case}] applied {result['applied']!r}, expected both edits")

    revisions = _revision_children(_body_paragraph(out, 1))
    expected = [
        ("ins", "Subject to Section 9, "),
        ("del", "sixty (60)"),
        ("ins", "thirty (30)"),
    ]
    if revisions != expected:
        failures.append(f"[{case}] revision sequence {revisions!r} != {expected!r}")

    with tempfile.TemporaryDirectory() as tmp:
        doc, path = _reopen(out, Path(tmp), "accept")
        try:
            doc.accept_all()
            doc.save()
        finally:
            doc.close()
        accepted = "".join(
            t.text or "" for t in _body_paragraph(path.read_bytes(), 1).iter(_qn("t"))
        )
        want = "  Subject to Section 9, The Term shall be thirty (30) days.  "
        if accepted != want:
            failures.append(f"[{case}] accept_all text {accepted!r} != {want!r}")


# ---------------------------------------------------------------------------
# Fix round 1, finding 3: the footnote anchor is the issue's first <w:ins>
# in DOCUMENT ORDER, not the first one recorded
# ---------------------------------------------------------------------------


def test_footnote_anchors_on_the_first_ins_in_document_order(failures: list) -> None:
    case = "footnote_anchors_on_the_first_ins_in_document_order"
    docx_bytes = _make_docx([_TWO_EDIT_TEXT])
    # ONE issue, TWO edits in one paragraph: pass 1 applies them in
    # descending-offset order, so the issue's revision ids are RECORDED
    # last-first.
    proven, _ = _prove(
        docx_bytes,
        [
            {"op": "keep", "text": "The Term shall be "},
            {"op": "delete", "text": "sixty (60)", "issue_key": "TERM-1"},
            {"op": "insert", "text": "thirty (30)", "issue_key": "TERM-1"},
            {"op": "keep", "text": " days and the Notice Period shall be "},
            {"op": "delete", "text": "ten (10)", "issue_key": "TERM-1"},
            {"op": "insert", "text": "five (5)", "issue_key": "TERM-1"},
            {"op": "keep", "text": " days."},
        ],
    )
    if proven["status"] != "proven":
        failures.append(f"[{case}] fixture transcript did not prove: {proven['failures']!r}")
        return

    rationale = "Both figures aligned with the negotiated position."
    result = redline_block_apply.apply_block_transcript(
        docx_bytes,
        proven,
        author=AUTHOR,
        timestamp_iso=TIMESTAMP,
        rationale_by_issue={"TERM-1": rationale},
    )
    out = result["docx_bytes"]
    if not out:
        failures.append(f"[{case}] no docx_bytes returned")
        return

    paragraph = _body_paragraph(out, 1)
    anchors = [
        "".join(t.text or "" for t in ins.iter(_qn("t")))
        for ins in paragraph.iter(_qn("ins"))
        if list(ins.iter(_qn("footnoteReference")))
    ]
    # "thirty (30)" is the FIRST insertion a reader reaches; "five (5)" is
    # the one recorded first, because pass 1 applies descending offsets.
    if anchors != ["thirty (30)"]:
        failures.append(
            f"[{case}] the footnote reference hangs on {anchors!r}, not on the issue's "
            f"first <w:ins> in document order ('thirty (30)')"
        )


# ---------------------------------------------------------------------------
# Runner
# ---------------------------------------------------------------------------

TESTS = [
    test_two_edits_in_one_paragraph,
    test_pure_insertion_splits_run_and_keeps_rpr,
    test_occurrence_disambiguation_edits_the_second_match,
    test_both_occurrences_edited_needs_descending_order,
    test_replacement_and_pure_insertion_in_one_paragraph,
    test_insertions_at_both_paragraph_boundaries,
    test_delete_across_physical_paragraphs_fails_closed,
    test_footnote_body_is_itself_tracked,
    test_ooxml_util_is_the_single_definition,
    test_rejected_transcript_raises,
    test_edit_never_slides_onto_a_later_equal_paragraph,
    test_whitespace_padded_paragraph_still_applies,
    test_footnote_anchors_on_the_first_ins_in_document_order,
]


def main() -> int:
    if redline_block_apply is None:
        print(f"FAIL: {IMPORT_ERROR}")
        return 1

    failures: list = []
    for test in TESTS:
        before = len(failures)
        try:
            test(failures)
        except Exception as exc:  # noqa: BLE001
            failures.append(f"[{test.__name__}] raised {type(exc).__name__}: {exc}")
        if len(failures) == before:
            print(f"PASS: {test.__name__}")
        else:
            for f in failures[before:]:
                print(f"FAIL: {f}")

    print()
    if failures:
        print(f"FAIL: {len(failures)} issue(s) found.")
        return 1
    print("PASS: all redline_block_apply (issue #621) assertions satisfied.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
