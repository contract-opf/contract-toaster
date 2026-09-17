#!/usr/bin/env python3
"""
Slice test for issue #145: "every document with a first-page or second
header/footer fails projection proof 3 after both model passes -- docx-editor
rewrites header2.xml, which is outside the part allowlist".

## Root problem this proves fixed

Reproduced 4 of 4 times against the owner's real affiliation agreement, then
offline with no model at all: any document whose section wires a first-page,
even-page, or other non-default header/footer part (`word/header2.xml`,
`word/footer2.xml`, ... -- the "different first page" title-page layout Word
writes routinely, and any auto-described logo image writes into) fails
`redline_projections.verify_projections`' proof 3 (the part allowlist) even
when NOTHING in the transcript touches the header or footer at all. The
mechanism, pinned by the issue's own investigation: `scripts/redline_block_
apply.py`'s pass 1 opens the package with the pinned `docx-editor` and saves
it, and that save re-serializes EVERY XML part it reads -- for a header or
footer carrying an image with a multi-line alt-text description (`<wp:docPr
descr="...">`, which Word writes with escaped line breaks, `&#xA;`, for
every auto-described logo), the escaped line breaks come back as literal
spaces. That is a canonical change to a part outside `redline_projections.
DECLARED_REDLINE_PARTS`, which named only `header1.xml`/`footer1.xml` -- the
marker's own -- so every such document failed proof 3 after the model had
already been paid for twice, and no redline was delivered.

## What this file asserts

  1. A two-header/footer document -- a DEFAULT header/footer pair
     (`header1.xml`/`footer1.xml`, the marker's own) AND a first-page pair
     (`header2.xml`/`footer2.xml`), the first-page header carrying Word's
     own multi-line-alt-text shape -- verifies after a bare `docx_editor`
     open/save with NO block edits at all: this is the issue's own offline
     reproduction (its "Required verification" section), and the healthy
     result on the untouched tree is that it FAILS.
  2. The SAME document survives a REAL replacement edit through the actual
     compiled pipeline (`redline_block_apply.apply_block_transcript`) --
     pass 1's docx-editor open/save is what a lone no-op skips silently if
     there is no `replace`/`delete` edit to drive it (see `redline_block_
     apply.py`'s `replacement_work` gate), so this is the assertion that
     matters for a real review.
  3. GUARDRAIL, required by the issue ("do not widen the allowlist to every
     header/footer part wholesale: a real header edit must still fail"): a
     `word/header2.xml` mutated to carry different VISIBLE text -- not just
     the docx-editor whitespace fold -- still fails proof 3, and names the
     part. The narrow rule this issue adds tolerates exactly one artifact,
     not a free pass on the part.
  4. GUARDRAIL, the mirror: a compiled output with `word/header2.xml`
     dropped from the package entirely fails proof 3 too -- the part stays
     OUT of `ALLOWED_CHANGED_PARTS`, so its disappearance is a violation
     like any other part's would be.

Uses python-docx (test-only dependency, matching `tests/test_redline_
projections.py`) to wire the two header/footer pairs -- the one shape that
module's own dependency-free raw-OOXML builder does not need -- then raw
zipfile+string surgery to inject the alt-text image, same convention as
`tools/churn_docx.py`'s `first_page_header_footer` transform. Every fixture
is SYNTHETIC: no real document, no real party names.

Exit codes: 0 = pass, 1 = fail
"""

from __future__ import annotations

import io
import re
import sys
import tempfile
import zipfile
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parent.parent
SCRIPTS_DIR = REPO_ROOT / "scripts"
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))

import block_transcript  # noqa: E402
import docx_editor  # noqa: E402
import extraction_normalization_stage  # noqa: E402
import redline_block_apply  # noqa: E402
import redline_projections  # noqa: E402

AUTHOR = "contract-toaster"
TIMESTAMP = "2026-01-01T00:00:00Z"
DOCUMENT_PART = "word/document.xml"

_TERM_BODY = "The Term shall be sixty (60) days."
_FEES_BODY = "Fees are due in forty-five (45) days."

# Word's own shape for a multi-line auto-generated alt-text description --
# the ONLY part of this drawing issue #145's mechanism depends on.
_ALT_TEXT_DRAWING_PARAGRAPH = (
    "<w:p><w:r><w:drawing>"
    '<wp:inline xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing">'
    '<wp:extent cx="1308824" cy="404813"/>'
    '<wp:effectExtent b="0" l="0" r="0" t="0"/>'
    '<wp:docPr descr="Icon&#xA;&#xA;Description automatically generated" id="5" name="image1.png"/>'
    "<wp:cNvGraphicFramePr/>"
    "</wp:inline></w:drawing></w:r></w:p>"
)


def _build_two_header_docx() -> tuple[bytes, str]:
    """A minimal docx with a DEFAULT header/footer pair (`word/header1.xml`
    / `word/footer1.xml`) AND a first-page header/footer pair (`word/
    header2.xml` / `word/footer2.xml`) -- the "different first page" layout
    Word writes for a title page. Returns `(docx_bytes, first_page_header_part)`.

    The default header/footer is wired FIRST and the first-page one second,
    deliberately: python-docx names whichever header/footer part it creates
    first "header1.xml" regardless of its `w:type`, and asserting the
    first-page header landed on `header2.xml` (not `header1.xml`, which
    `redline_projections.DECLARED_REDLINE_PARTS` already allows to change
    wholesale) is what makes this fixture reproduce the issue rather than
    something the allowlist already tolerated.
    """
    import docx  # local import: python-docx is a test-only dependency

    document = docx.Document()
    document.add_paragraph("Section 1. Term", style="Heading 1")
    document.add_paragraph(_TERM_BODY)
    document.add_paragraph("Section 2. Fees", style="Heading 1")
    document.add_paragraph(_FEES_BODY)

    section = document.sections[0]
    section.header.paragraphs[0].add_run("Default header")
    section.footer.paragraphs[0].add_run("Default footer")
    section.different_first_page_header_footer = True
    section.first_page_header.paragraphs[0].add_run("First-page header")
    section.first_page_footer.paragraphs[0].add_run("First-page footer")

    buf = io.BytesIO()
    document.save(buf)
    wired_bytes = buf.getvalue()

    with zipfile.ZipFile(io.BytesIO(wired_bytes)) as zf:
        infos = zf.infolist()
        parts = {info.filename: zf.read(info.filename) for info in infos}

    doc_xml = parts[DOCUMENT_PART].decode("utf-8")
    header_ref = re.search(r'<w:headerReference\s+w:type="first"\s+r:id="(rId\d+)"\s*/>', doc_xml)
    if header_ref is None:
        raise AssertionError(
            "python-docx did not wire a first-type <w:headerReference> -- fixture is wrong"
        )
    header_rid = header_ref.group(1)

    rels_xml = parts["word/_rels/document.xml.rels"].decode("utf-8")
    rel_target = re.search(rf'<Relationship\s+Id="{header_rid}"[^>]*\bTarget="([^"]+)"', rels_xml)
    if rel_target is None:
        raise AssertionError(f"no relationship target for header id {header_rid!r} -- fixture is wrong")
    header_part = f"word/{rel_target.group(1)}"
    if header_part == "word/header1.xml":
        raise AssertionError(
            "the first-page header landed on header1.xml -- the fixture must wire the "
            "DEFAULT header/footer first so header1.xml stays the already-allowlisted part "
            "and header2.xml is the one this issue is actually about"
        )

    header_xml = parts[header_part].decode("utf-8")
    if "</w:hdr>" not in header_xml:
        raise AssertionError(f"{header_part!r} is not a <w:hdr> part -- fixture is wrong")
    parts[header_part] = header_xml.replace(
        "</w:hdr>", _ALT_TEXT_DRAWING_PARAGRAPH + "</w:hdr>", 1
    ).encode("utf-8")

    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zf_out:
        for info in infos:
            zf_out.writestr(info, parts[info.filename])
    return out.getvalue(), header_part


def _editor_roundtrip(docx_bytes: bytes) -> bytes:
    """Opens and saves `docx_bytes` through the pinned `docx-editor` with NO
    edits at all -- the issue's own offline reproduction."""
    with tempfile.TemporaryDirectory(prefix="projection-allowlist-145-") as tmp:
        tmp_path = Path(tmp)
        input_path = tmp_path / "in.docx"
        input_path.write_bytes(docx_bytes)
        doc = docx_editor.Document.open(
            input_path, author=AUTHOR, workspace_dir=str(tmp_path / "ws")
        )
        doc.save()
        doc.close()
        return input_path.read_bytes()


def _repack(docx_bytes: bytes, *, replace=None, remove=None) -> bytes:
    """`docx_bytes` with `replace = {part: new_bytes}` swapped in and every
    name in `remove` dropped."""
    replace = replace or {}
    remove = set(remove or ())
    with zipfile.ZipFile(io.BytesIO(docx_bytes)) as zf:
        infos = zf.infolist()
        originals = {info.filename: zf.read(info.filename) for info in infos}
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf_out:
        for info in infos:
            if info.filename in remove:
                continue
            zf_out.writestr(info, replace.get(info.filename, originals[info.filename]))
    return buf.getvalue()


# ---------------------------------------------------------------------------


def test_untouched_two_header_document_verifies_after_editor_roundtrip(failures: list) -> None:
    """The issue's own "Required verification": a bare `docx_editor`
    open/save on a two-header/footer document, with NO block edits at all,
    must still verify -- proof 3 must not fail-closed on an artifact the
    editor's own save introduces regardless of what a review asked for."""
    case = "untouched_two_header_document_verifies_after_editor_roundtrip"
    docx_bytes, header_part = _build_two_header_docx()
    if header_part != "word/header2.xml":
        failures.append(f"[{case}] fixture wired the first-page header onto {header_part!r}, not header2.xml")
        return

    output_bytes = _editor_roundtrip(docx_bytes)
    with zipfile.ZipFile(io.BytesIO(docx_bytes)) as zf:
        source_header2 = zf.read(header_part)
    with zipfile.ZipFile(io.BytesIO(output_bytes)) as zf:
        output_header2 = zf.read(header_part)
    if source_header2 == output_header2:
        failures.append(
            f"[{case}] docx-editor did not actually change {header_part!r} -- this fixture "
            "no longer reproduces the artifact and proves nothing"
        )
        return

    report = redline_projections.verify_projections(
        docx_bytes, output_bytes, {"block_ops": []}, revision_ids=set(), applied_edits=[]
    )
    if report["status"] != "verified":
        failures.append(f"[{case}] a bare editor round-trip failed the proofs: {report['failures']!r}")


def test_replace_edit_compiles_through_the_real_writer(failures: list) -> None:
    """The property that matters for a real review: a REPLACEMENT edit --
    which drives pass 1's docx-editor open/save, unlike a pure insertion --
    on a document carrying the first-page-header artifact must still
    compile through `apply_block_transcript`."""
    case = "replace_edit_compiles_through_the_real_writer"
    docx_bytes, _header_part = _build_two_header_docx()

    norm = extraction_normalization_stage.extract_and_normalize(docx_bytes)
    if norm["status"] != "normalized":
        failures.append(f"[{case}] fixture did not normalize: status={norm.get('status')!r}")
        return
    block_map = extraction_normalization_stage.build_block_map(norm["paragraphs"])
    term_id = next((bid for bid, block in block_map.items() if _TERM_BODY in block["text"]), None)
    if term_id is None:
        failures.append(f"[{case}] could not find the Term block")
        return

    proven = block_transcript.validate_block_patches(
        [
            {
                "block_id": term_id,
                "segments": [
                    {"op": "keep", "text": "The Term shall be "},
                    {"op": "delete", "text": "sixty (60)", "issue_key": "TERM-1"},
                    {"op": "insert", "text": "thirty (30)", "issue_key": "TERM-1"},
                    {"op": "keep", "text": " days."},
                ],
            }
        ],
        [],
        block_map,
    )
    if proven["status"] != "proven":
        failures.append(f"[{case}] the replacement did not prove: {[f.get('reason') for f in proven['failures']]}")
        return

    result = redline_block_apply.apply_block_transcript(
        docx_bytes, proven, author=AUTHOR, timestamp_iso=TIMESTAMP
    )
    if result["docx_bytes"] is None or result["failures"]:
        failures.append(
            f"[{case}] a replacement edit on a two-header/footer document did not compile: "
            f"applied={len(result['applied'])} failures={result['failures']!r}"
        )


def test_a_real_header_edit_still_fails_the_allowlist(failures: list) -> None:
    """GUARDRAIL the issue requires explicitly: the narrow rule tolerates
    ONLY the docx-editor whitespace fold, not a free pass on `header2.xml`
    -- a redline that actually rewrites the first-page header's VISIBLE
    text must still fail proof 3."""
    case = "a_real_header_edit_still_fails_the_allowlist"
    docx_bytes, header_part = _build_two_header_docx()
    output_bytes = _editor_roundtrip(docx_bytes)

    with zipfile.ZipFile(io.BytesIO(output_bytes)) as zf:
        header_xml = zf.read(header_part).decode("utf-8")
    if "First-page header" not in header_xml:
        failures.append(f"[{case}] fixture is wrong -- {header_part!r} does not carry its own visible text")
        return
    tampered_xml = header_xml.replace("First-page header", "Rewritten first-page header")
    tampered_bytes = _repack(output_bytes, replace={header_part: tampered_xml.encode("utf-8")})

    report = redline_projections.verify_projections(
        docx_bytes, tampered_bytes, {"block_ops": []}, revision_ids=set(), applied_edits=[]
    )
    allowlist_parts = [
        entry.get("part")
        for entry in report["failures"]
        if entry["proof"] == redline_projections.PROOF_PART_ALLOWLIST
    ]
    if header_part not in allowlist_parts:
        failures.append(
            f"[{case}] a real content edit to {header_part!r} passed the allowlist proof: "
            f"{report['failures']!r}"
        )


def test_a_dropped_header_part_still_fails_the_allowlist(failures: list) -> None:
    """GUARDRAIL, the mirror: `word/header2.xml` disappearing from the
    package entirely -- a writer bug that silently drops it -- must fail
    proof 3 too. The part stays OUT of `ALLOWED_CHANGED_PARTS`, so its
    disappearance is a violation like any other part's would be."""
    case = "a_dropped_header_part_still_fails_the_allowlist"
    docx_bytes, header_part = _build_two_header_docx()
    output_bytes = _editor_roundtrip(docx_bytes)
    dropped_bytes = _repack(output_bytes, remove=[header_part])

    report = redline_projections.verify_projections(
        docx_bytes, dropped_bytes, {"block_ops": []}, revision_ids=set(), applied_edits=[]
    )
    allowlist_parts = [
        entry.get("part")
        for entry in report["failures"]
        if entry["proof"] == redline_projections.PROOF_PART_ALLOWLIST
    ]
    if header_part not in allowlist_parts:
        failures.append(
            f"[{case}] a dropped {header_part!r} passed the allowlist proof: {report['failures']!r}"
        )


# ---------------------------------------------------------------------------
# Runner
# ---------------------------------------------------------------------------

TESTS = [
    test_untouched_two_header_document_verifies_after_editor_roundtrip,
    test_replace_edit_compiles_through_the_real_writer,
    test_a_real_header_edit_still_fails_the_allowlist,
    test_a_dropped_header_part_still_fails_the_allowlist,
]


def main() -> int:
    failures: list = []
    for test in TESTS:
        before = len(failures)
        try:
            test(failures)
        except Exception as exc:  # noqa: BLE001
            failures.append(f"[{test.__name__}] raised {type(exc).__name__}: {exc}")
        print(("PASS: " if len(failures) == before else "FAIL: ") + test.__name__)

    print()
    if failures:
        for f in failures:
            print(f"  - {f}")
        print(f"\nFAIL: {len(failures)} issue(s) found.")
        return 1
    print("PASS: all projection-allowlist multi-header (issue #145) assertions satisfied.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
