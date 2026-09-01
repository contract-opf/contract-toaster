#!/usr/bin/env python3
"""
Slice test (TDD) for issue #622: "whole-block ops -- delete_block and
insert_block_after".

## Root problem this proves fixed

Before this slice the block compiler (`scripts/redline_block_apply.py`,
issue #621) could only write edits that named a sub-block SPAN. An issue
with no in-place span to edit -- "this clause must go", "this agreement
needs an audit-rights clause it does not have" -- had nowhere to land:
`proven["block_ops"]` was handed back untouched under `deferred_block_ops`
and every such issue degraded to flag-only. This file drives the two ops
that close that gap through the REAL validator
(`block_transcript.validate_block_patches`, issue #620), so every op the
compiler consumes here was proven against the fixture's own bytes -- no
hand-built op shapes, no offsets written by the test.

## What this test asserts (mirrors the issue's Acceptance criteria)

  1. AC 1 -- an `insert_block_after` paragraph appears immediately after its
     anchor block's LAST physical `<w:p>`, with tracked-insert semantics
     (`<w:ins>` carrying `w:id`/`w:author`/`w:date`, AND the new paragraph's
     own mark marked inserted so rejecting it leaves no stray pilcrow), and
     its issue's footnoted rationale is attached INSIDE that `<w:ins>`.
  2. `anchor_block_id == "start"` inserts BEFORE the document's first block.
  3. A block that joins several sibling `<w:p>`s is anchored on its LAST
     one, so the new clause follows the block instead of landing inside it.
  4. Two insertions naming one anchor keep their transcript order.
  5. AC 2 -- a `delete_block` renders as a full-paragraph tracked deletion:
     every run wrapped in `<w:del>` with `<w:delText>` (never `<w:t>`, the
     incorrect shortcut `scripts/redline_inplace.py` documents) and the
     paragraph MARK marked deleted; reject-all restores the original
     paragraph structure exactly in text space, and accept-all removes the
     clause's text.
  6. A `delete_block` covers EVERY physical `<w:p>` of a multi-paragraph
     (list) block, not just the first ...
  7. ... and every run those paragraphs SHOW, including ones nested in a
     counterparty's pending `<w:ins>` or in a `<w:fldSimple>` field result.
     Text surviving inside a paragraph whose mark was deleted, reported as
     applied, is the worst outcome this op can produce.
  8. AC 3 -- a block op aimed at a table-cell block is a structured
     `block_op_unsupported_in_table` failure naming the op and the block id,
     and it does not block the other ops in the same transcript.
  9. ... while a within-cell SPAN edit still applies: the table rule is
     scoped to whole-block ops, exactly as the issue's Scope says.
 10. Block ops and span edits in ONE transcript all land, and the `<w:p>`
     renumbering an insertion causes never corrupts a sibling op's target.
 11. Round-trip verification passes on every produced document, and
     `docx-editor` itself reopens the output and cleanly
     `accept_all()`/`reject_all()`s it.

Fixtures are built with python-docx (a test-only dependency, matching
`tests/test_redline_block_apply.py`), plus a dependency-free raw-OOXML
builder for the two shapes python-docx cannot author -- `<w:ins>` and
`<w:fldSimple>` -- the same convention this repo's other OOXML fixture
builders use.
All SYNTHETIC: never a real document, never real party names.

Exit codes: 0 = pass, 1 = fail
"""

from __future__ import annotations

import io
import sys
import tempfile
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parent.parent
SCRIPTS_DIR = REPO_ROOT / "scripts"
sys.path.insert(0, str(SCRIPTS_DIR))

import block_transcript  # noqa: E402
import docx_editor  # noqa: E402
import extraction_normalization_stage  # noqa: E402
import redline_block_apply  # noqa: E402
import redline_generate  # noqa: E402

WORD_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

AUTHOR = "contract-toaster"
TIMESTAMP = "2026-01-01T00:00:00Z"


def _qn(tag: str) -> str:
    return f"{{{WORD_NS}}}{tag}"


# ---------------------------------------------------------------------------
# Synthetic fixture construction (python-docx allowed in tests only)
#
# Each heading paragraph carries a real Heading style so
# `clause_boundaries.is_boundary_paragraph_ooxml` starts a logical block at
# it; the body paragraphs under it are that block's PHYSICAL paragraphs.
# Nothing here is drawn from any real document.
# ---------------------------------------------------------------------------


def _make_sectioned_docx(sections: list, *, style: str = None) -> bytes:
    """`[(heading, [body, ...]), ...]` -> docx bytes."""
    import docx  # local import: python-docx is a test-only dependency

    document = docx.Document()
    for heading, bodies in sections:
        document.add_paragraph(heading, style="Heading 1")
        for body in bodies:
            if style:
                document.add_paragraph(body, style=style)
            else:
                paragraph = document.add_paragraph()
                paragraph.add_run(body)
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def _make_table_docx() -> bytes:
    """A body clause, then a heading whose only content is a one-cell table
    -- so the cell's paragraph is a logical block of its OWN, which is what
    a block op aimed at a table cell looks like."""
    import docx  # local import: python-docx is a test-only dependency

    document = docx.Document()
    document.add_paragraph("Section 1. Term", style="Heading 1")
    paragraph = document.add_paragraph()
    paragraph.add_run(_TERM_BODY)
    document.add_paragraph("Section 2. Schedule", style="Heading 1")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).paragraphs[0].add_run(_CELL_BODY)
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


# A minimal, dependency-free OOXML package, for the fixtures whose markup
# python-docx cannot author (`<w:ins>`, `<w:fldSimple>`) -- the same
# convention `tests/test_extraction_normalization_stage_80.py` uses. Still synthetic:
# every string below is invented for this test.
_CONTENT_TYPES_XML = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
    '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
    '<Default Extension="rels" '
    'ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
    '<Default Extension="xml" ContentType="application/xml"/>'
    '<Override PartName="/word/document.xml" '
    'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.'
    'document.main+xml"/>'
    "</Types>"
)

_RELS_XML = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
    '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
    '<Relationship Id="rId1" '
    'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/'
    'officeDocument" Target="word/document.xml"/>'
    "</Relationships>"
)


def _build_raw_docx(body_paragraphs_xml: str) -> bytes:
    document_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        f'<w:document xmlns:w="{WORD_NS}">'
        f"<w:body>{body_paragraphs_xml}<w:sectPr/></w:body>"
        "</w:document>"
    )
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("[Content_Types].xml", _CONTENT_TYPES_XML)
        zf.writestr("_rels/.rels", _RELS_XML)
        zf.writestr("word/document.xml", document_xml)
    return buf.getvalue()


def _raw_heading_p(text: str) -> str:
    return (
        '<w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr>'
        f"<w:r><w:t>{text}</w:t></w:r></w:p>"
    )


def _prove(docx_bytes: bytes, *, patches=None, ops=None):
    """Run the REAL #620 validator over this transcript against the
    fixture's own block map, so every op the compiler consumes was proven --
    never hand-written."""
    norm = extraction_normalization_stage.extract_and_normalize(docx_bytes)
    assert norm["status"] == "normalized", norm
    block_map = extraction_normalization_stage.build_block_map(norm["paragraphs"])
    proven = block_transcript.validate_block_patches(
        list(patches or []), list(ops or []), block_map
    )
    return proven, block_map


# ---------------------------------------------------------------------------
# Readers
# ---------------------------------------------------------------------------


def _document_root(docx_bytes: bytes) -> ET.Element:
    with zipfile.ZipFile(io.BytesIO(docx_bytes)) as zf:
        return ET.fromstring(zf.read("word/document.xml"))


def _footnotes_root(docx_bytes: bytes):
    with zipfile.ZipFile(io.BytesIO(docx_bytes)) as zf:
        if "word/footnotes.xml" not in zf.namelist():
            return None
        return ET.fromstring(zf.read("word/footnotes.xml"))


def _paragraphs(docx_bytes: bytes) -> list:
    return list(_document_root(docx_bytes).iter(_qn("p")))


def _paragraph_texts(docx_bytes: bytes) -> list:
    """One entry per `<w:p>`, in document order: every `<w:t>` AND
    `<w:delText>` it carries. This is the "text space" the issue's
    reject-all criterion is stated in."""
    out = []
    for p in _paragraphs(docx_bytes):
        out.append(
            "".join(
                el.text or ""
                for el in p.iter()
                if el.tag in (_qn("t"), _qn("delText"))
            )
        )
    return out


def _visible_paragraph_texts(docx_bytes: bytes) -> list:
    """One entry per `<w:p>`: only `<w:t>` -- i.e. what survives once a
    deletion's `<w:delText>` is excluded, which is the accepted view."""
    return [
        "".join(t.text or "" for t in p.iter(_qn("t"))) for p in _paragraphs(docx_bytes)
    ]


def _paragraph_mark_revision(p: ET.Element, tag: str):
    """The `<w:pPr><w:rPr><w:ins|w:del>` marker on `p`'s paragraph mark, or
    None."""
    pPr = p.find(_qn("pPr"))
    if pPr is None:
        return None
    rpr = pPr.find(_qn("rPr"))
    if rpr is None:
        return None
    return rpr.find(_qn(tag))


def _reopen(docx_bytes: bytes, tmp_dir: Path, name: str):
    work = tmp_dir / name
    work.mkdir()
    path = work / "doc.docx"
    path.write_bytes(docx_bytes)
    doc = docx_editor.Document.open(path, author=AUTHOR, workspace_dir=str(work / "ws"))
    return doc, path


def _resolved(docx_bytes: bytes, tmp_dir: Path, name: str, how: str) -> bytes:
    doc, path = _reopen(docx_bytes, tmp_dir, name)
    try:
        getattr(doc, how)()
        doc.save()
    finally:
        doc.close()
    return path.read_bytes()


def _round_trips(case: str, out: bytes, failures: list) -> None:
    try:
        redline_generate.verify_docx_round_trip(out)
    except ValueError as exc:
        failures.append(f"[{case}] verify_docx_round_trip raised: {exc}")


# ---------------------------------------------------------------------------
# Fixture text (synthetic)
# ---------------------------------------------------------------------------

_TERM_BODY = "The Term shall be sixty (60) days."
_FEES_BODY = "Fees are due in thirty (30) days."
_RENEWAL_BODY = "The Agreement renews automatically."
_CELL_BODY = "Deliverables are accepted on receipt."
_NEW_CLAUSE = "Each party shall retain audit rights for two (2) years."


# ---------------------------------------------------------------------------
# AC 1: insert_block_after
# ---------------------------------------------------------------------------


def test_insert_block_after_anchor(failures: list) -> None:
    case = "insert_block_after_anchor"
    docx_bytes = _make_sectioned_docx(
        [("Section 1. Term", [_TERM_BODY]), ("Section 2. Fees", [_FEES_BODY])]
    )
    before = _paragraph_texts(docx_bytes)
    rationale = "An audit right is required by the negotiating position."
    proven, _ = _prove(
        docx_bytes,
        ops=[
            {
                "op": "insert_block_after",
                "anchor_block_id": "p0001",
                "new_text": _NEW_CLAUSE,
                "issue_key": "AUDIT-1",
            }
        ],
    )
    if proven["status"] != "proven":
        failures.append(f"[{case}] fixture transcript did not prove: {proven['failures']!r}")
        return

    result = redline_block_apply.apply_block_transcript(
        docx_bytes,
        proven,
        author=AUTHOR,
        timestamp_iso=TIMESTAMP,
        rationale_by_issue={"AUDIT-1": rationale},
    )
    if result["failures"]:
        failures.append(f"[{case}] unexpected failures: {result['failures']!r}")
    out = result["docx_bytes"]
    if not out:
        failures.append(f"[{case}] no docx_bytes returned -- the block op was not written")
        return
    _round_trips(case, out, failures)

    if [entry["kind"] for entry in result["applied"]] != ["insert_block_after"]:
        failures.append(f"[{case}] applied entries {result['applied']!r}")

    # AC 1a: the new paragraph sits immediately after the anchor block's
    # last physical <w:p> (0 = "Section 1. Term", 1 = its body).
    texts = _paragraph_texts(out)
    expected = before[:2] + [_NEW_CLAUSE] + before[2:]
    if texts != expected:
        failures.append(f"[{case}] paragraph layout {texts!r} != {expected!r}")
        return

    new_p = _paragraphs(out)[2]

    # AC 1b: tracked-insert semantics -- the body run is inside a <w:ins>
    # carrying all three revision attributes.
    body_ins = [
        ins
        for ins in new_p.iter(_qn("ins"))
        if "".join(t.text or "" for t in ins.iter(_qn("t")))
    ]
    if len(body_ins) != 1:
        failures.append(f"[{case}] expected exactly one text-bearing <w:ins>, got {len(body_ins)}")
        return
    ins = body_ins[0]
    for attr, want in (("author", AUTHOR), ("date", TIMESTAMP)):
        if ins.get(_qn(attr)) != want:
            failures.append(f"[{case}] <w:ins> w:{attr} is {ins.get(_qn(attr))!r}, want {want!r}")
    if ins.get(_qn("id")) in (None, ""):
        failures.append(f"[{case}] <w:ins> carries no w:id")
    if list(new_p.iter(_qn("del"))):
        failures.append(f"[{case}] an insertion emitted a <w:del>")

    # AC 1c: the new paragraph's own MARK is marked inserted, so rejecting
    # the change leaves no stray pilcrow behind.
    mark = _paragraph_mark_revision(new_p, "ins")
    if mark is None:
        failures.append(
            f"[{case}] the inserted paragraph's mark is NOT tracked "
            f"(<w:pPr><w:rPr><w:ins/></w:rPr></w:pPr> missing)"
        )
    elif mark.get(_qn("author")) != AUTHOR or mark.get(_qn("date")) != TIMESTAMP:
        failures.append(f"[{case}] paragraph-mark <w:ins> author/date not stamped")

    # AC 1d: the issue's footnote is attached, and INSIDE the body <w:ins>
    # -- never on the paragraph-mark marker, which holds no run content.
    refs = [r.get(_qn("id")) for r in ins.iter(_qn("footnoteReference"))]
    if len(refs) != 1:
        failures.append(f"[{case}] expected 1 footnoteReference inside the <w:ins>, got {refs!r}")
    else:
        footnotes_root = _footnotes_root(out)
        if footnotes_root is None:
            failures.append(f"[{case}] no word/footnotes.xml in the output")
        else:
            body = ""
            for fn in footnotes_root.findall(_qn("footnote")):
                if fn.get(_qn("id")) == refs[0]:
                    body = "".join(t.text or "" for t in fn.iter(_qn("t")))
            if rationale not in body:
                failures.append(f"[{case}] the inserted clause's footnote body is {body!r}")
    if _paragraph_mark_revision(new_p, "ins") is not None and list(
        _paragraph_mark_revision(new_p, "ins").iter(_qn("footnoteReference"))
    ):
        failures.append(f"[{case}] a footnote reference was written into <w:pPr><w:rPr>")

    with tempfile.TemporaryDirectory() as tmp:
        tmp_path = Path(tmp)
        accepted = _visible_paragraph_texts(_resolved(out, tmp_path, "accept", "accept_all"))
        want = before[:2] + [_NEW_CLAUSE] + before[2:]
        if accepted != want:
            failures.append(f"[{case}] accept_all layout {accepted!r} != {want!r}")
        # Reject-all removes the inserted text entirely. The now-empty <w:p>
        # may survive (docx-editor removes revisions, never paragraphs), so
        # the assertion is on the text: nothing of the new clause is left.
        rejected = [
            text
            for text in _visible_paragraph_texts(_resolved(out, tmp_path, "reject", "reject_all"))
            if text
        ]
        if rejected != [text for text in before if text]:
            failures.append(f"[{case}] reject_all text {rejected!r} != {before!r}")


def test_insert_block_at_start(failures: list) -> None:
    case = "insert_block_at_start"
    docx_bytes = _make_sectioned_docx([("Section 1. Term", [_TERM_BODY])])
    before = _paragraph_texts(docx_bytes)
    proven, _ = _prove(
        docx_bytes,
        ops=[
            {
                "op": "insert_block_after",
                "anchor_block_id": block_transcript.ANCHOR_START,
                "new_text": _NEW_CLAUSE,
                "issue_key": "AUDIT-1",
            }
        ],
    )
    if proven["status"] != "proven":
        failures.append(f"[{case}] fixture transcript did not prove: {proven['failures']!r}")
        return

    result = redline_block_apply.apply_block_transcript(
        docx_bytes, proven, author=AUTHOR, timestamp_iso=TIMESTAMP
    )
    if result["failures"]:
        failures.append(f"[{case}] unexpected failures: {result['failures']!r}")
    out = result["docx_bytes"]
    if not out:
        failures.append(f"[{case}] no docx_bytes returned")
        return
    _round_trips(case, out, failures)

    # `"start"` anchors on no block: the new paragraph goes BEFORE the first
    # block's first physical <w:p>, i.e. between the heading (<w:p> 0) and
    # the body it introduces (<w:p> 1).
    texts = _paragraph_texts(out)
    expected = before[:1] + [_NEW_CLAUSE] + before[1:]
    if texts != expected:
        failures.append(f"[{case}] paragraph layout {texts!r} != {expected!r}")


def test_insert_lands_after_the_anchors_last_physical_paragraph(failures: list) -> None:
    """A logical block joins several sibling `<w:p>`s. "After the block"
    means after the LAST of them -- landing after the first would drop the
    new clause into the MIDDLE of the block it was meant to follow."""
    case = "insert_lands_after_the_anchors_last_physical_paragraph"
    clauses = ["First list clause of the block.", "Second list clause of the block."]
    docx_bytes = _make_sectioned_docx([("Section 1. Term", clauses)], style="List Number")
    before = _paragraph_texts(docx_bytes)
    proven, block_map = _prove(
        docx_bytes,
        ops=[
            {
                "op": "insert_block_after",
                "anchor_block_id": "p0001",
                "new_text": _NEW_CLAUSE,
                "issue_key": "AUDIT-1",
            }
        ],
    )
    if proven["status"] != "proven":
        failures.append(f"[{case}] fixture transcript did not prove: {proven['failures']!r}")
        return
    if len(block_map["p0001"]["physical_spans"]) != 2:
        failures.append(
            f"[{case}] the fixture anchor block is not multi-paragraph: "
            f"{block_map['p0001']['physical_spans']!r}"
        )
        return

    result = redline_block_apply.apply_block_transcript(
        docx_bytes, proven, author=AUTHOR, timestamp_iso=TIMESTAMP
    )
    out = result["docx_bytes"]
    if result["failures"] or not out:
        failures.append(f"[{case}] failures={result['failures']!r} bytes={bool(out)}")
        return
    _round_trips(case, out, failures)

    texts = _paragraph_texts(out)
    expected = before + [_NEW_CLAUSE]
    if texts != expected:
        failures.append(
            f"[{case}] paragraph layout {texts!r} != {expected!r} -- the insertion did not "
            f"land after the anchor block's LAST physical <w:p>"
        )


def test_two_inserts_after_one_anchor_keep_transcript_order(failures: list) -> None:
    case = "two_inserts_after_one_anchor_keep_transcript_order"
    docx_bytes = _make_sectioned_docx([("Section 1. Term", [_TERM_BODY])])
    before = _paragraph_texts(docx_bytes)
    first = "Each party shall retain audit rights for two (2) years."
    second = "Audit costs are borne by the requesting party."
    proven, _ = _prove(
        docx_bytes,
        ops=[
            {
                "op": "insert_block_after",
                "anchor_block_id": "p0001",
                "new_text": first,
                "issue_key": "AUDIT-1",
            },
            {
                "op": "insert_block_after",
                "anchor_block_id": "p0001",
                "new_text": second,
                "issue_key": "AUDIT-2",
            },
        ],
    )
    if proven["status"] != "proven":
        failures.append(f"[{case}] fixture transcript did not prove: {proven['failures']!r}")
        return

    result = redline_block_apply.apply_block_transcript(
        docx_bytes, proven, author=AUTHOR, timestamp_iso=TIMESTAMP
    )
    out = result["docx_bytes"]
    if result["failures"] or not out:
        failures.append(f"[{case}] failures={result['failures']!r} bytes={bool(out)}")
        return
    _round_trips(case, out, failures)

    texts = _paragraph_texts(out)
    expected = before[:2] + [first, second] + before[2:]
    if texts != expected:
        failures.append(
            f"[{case}] two insertions after one anchor came out {texts!r} != {expected!r} "
            f"-- the second anchored on the original paragraph instead of the first "
            f"insertion, which reverses them"
        )


# ---------------------------------------------------------------------------
# AC 2: delete_block
# ---------------------------------------------------------------------------


def test_delete_block_deletes_runs_and_paragraph_mark(failures: list) -> None:
    case = "delete_block_deletes_runs_and_paragraph_mark"
    docx_bytes = _make_sectioned_docx(
        [("Section 1. Term", [_TERM_BODY]), ("Section 2. Renewal", [_RENEWAL_BODY])]
    )
    before = _paragraph_texts(docx_bytes)
    proven, _ = _prove(
        docx_bytes,
        ops=[{"op": "delete_block", "block_id": "p0002", "issue_key": "RENEW-1"}],
    )
    if proven["status"] != "proven":
        failures.append(f"[{case}] fixture transcript did not prove: {proven['failures']!r}")
        return

    result = redline_block_apply.apply_block_transcript(
        docx_bytes, proven, author=AUTHOR, timestamp_iso=TIMESTAMP
    )
    if result["failures"]:
        failures.append(f"[{case}] unexpected failures: {result['failures']!r}")
    out = result["docx_bytes"]
    if not out:
        failures.append(f"[{case}] no docx_bytes returned -- the block op was not written")
        return
    _round_trips(case, out, failures)

    if [entry["kind"] for entry in result["applied"]] != ["delete_block"]:
        failures.append(f"[{case}] applied entries {result['applied']!r}")

    target = _paragraphs(out)[3]  # 0 h1, 1 body1, 2 h2, 3 body2

    # Every run is inside a <w:del> ...
    for run in target.iter(_qn("r")):
        if not any(run in list(d.iter(_qn("r"))) for d in target.iter(_qn("del"))):
            failures.append(f"[{case}] a run of the deleted block is not inside a <w:del>")
            break
    # ... using <w:delText>, never <w:t>: `<w:t>` inside a `<w:del>` renders
    # wrong in Word's Reviewing pane (scripts/redline_inplace.py, "Rewrite").
    if list(target.iter(_qn("t"))):
        failures.append(
            f"[{case}] the deleted paragraph still carries <w:t> -- deleted text must be "
            f"<w:delText>"
        )
    if not list(target.iter(_qn("delText"))):
        failures.append(f"[{case}] the deleted paragraph carries no <w:delText> at all")
    for del_el in target.iter(_qn("del")):
        if del_el.get(_qn("author")) != AUTHOR or del_el.get(_qn("date")) != TIMESTAMP:
            failures.append(f"[{case}] a <w:del> is not stamped with this call's author/date")
            break

    # The paragraph MARK is deleted too -- the difference between deleting a
    # paragraph and emptying one.
    mark = _paragraph_mark_revision(target, "del")
    if mark is None:
        failures.append(
            f"[{case}] the paragraph MARK is not deleted "
            f"(<w:pPr><w:rPr><w:del/></w:rPr></w:pPr> missing) -- accepting the change "
            f"would leave a blank paragraph behind"
        )
    elif mark.get(_qn("id")) in (None, ""):
        failures.append(f"[{case}] the paragraph-mark <w:del> carries no w:id")

    # The OTHER section is untouched.
    if list(_paragraphs(out)[1].iter(_qn("del"))):
        failures.append(f"[{case}] the deletion contaminated Section 1's paragraph")

    with tempfile.TemporaryDirectory() as tmp:
        tmp_path = Path(tmp)
        # AC 2: rejecting all changes restores the original paragraph
        # structure exactly, in text space.
        rejected = _paragraph_texts(_resolved(out, tmp_path, "reject", "reject_all"))
        if rejected != before:
            failures.append(f"[{case}] reject_all text {rejected!r} != {before!r}")
        # ... and accepting removes the clause's text.
        accepted = _visible_paragraph_texts(_resolved(out, tmp_path, "accept", "accept_all"))
        if _RENEWAL_BODY in "".join(accepted):
            failures.append(f"[{case}] accept_all left the deleted clause behind: {accepted!r}")
        if _TERM_BODY not in "".join(accepted):
            failures.append(f"[{case}] accept_all also removed the untouched clause")


def test_delete_block_covers_every_physical_paragraph(failures: list) -> None:
    case = "delete_block_covers_every_physical_paragraph"
    clauses = ["First list clause of the block.", "Second list clause of the block."]
    docx_bytes = _make_sectioned_docx(
        [("Section 1. Term", clauses)], style="List Number"
    )
    before = _paragraph_texts(docx_bytes)
    proven, block_map = _prove(
        docx_bytes,
        ops=[{"op": "delete_block", "block_id": "p0001", "issue_key": "TERM-1"}],
    )
    if proven["status"] != "proven":
        failures.append(f"[{case}] fixture transcript did not prove: {proven['failures']!r}")
        return
    if len(block_map["p0001"]["physical_spans"]) != 2:
        failures.append(
            f"[{case}] the fixture block is not multi-paragraph: "
            f"{block_map['p0001']['physical_spans']!r}"
        )
        return

    result = redline_block_apply.apply_block_transcript(
        docx_bytes, proven, author=AUTHOR, timestamp_iso=TIMESTAMP
    )
    if result["failures"]:
        failures.append(f"[{case}] unexpected failures: {result['failures']!r}")
    out = result["docx_bytes"]
    if not out:
        failures.append(f"[{case}] no docx_bytes returned")
        return
    _round_trips(case, out, failures)

    for index in (1, 2):  # both physical paragraphs of the one logical block
        target = _paragraphs(out)[index]
        if not list(target.iter(_qn("delText"))):
            failures.append(
                f"[{case}] physical paragraph {index} of the block was NOT deleted -- "
                f"a whole-block delete must cover every <w:p> the block joins"
            )
        if _paragraph_mark_revision(target, "del") is None:
            failures.append(f"[{case}] physical paragraph {index} kept an undeleted mark")

    with tempfile.TemporaryDirectory() as tmp:
        rejected = _paragraph_texts(_resolved(out, Path(tmp), "reject", "reject_all"))
        if rejected != before:
            failures.append(f"[{case}] reject_all text {rejected!r} != {before!r}")


def test_delete_block_reaches_runs_nested_in_containers(failures: list) -> None:
    """A clause whose runs are not all direct `<w:p>` children.

    Both shapes here are ordinary in counterparty documents: a pending
    `<w:ins>` (the other side edited with track changes on -- the extractor's
    accept-all disposition makes its text part of the block's operative
    text) and a `<w:fldSimple>` field result. If the deletion only wrapped
    top-level `<w:r>`s, that text would SURVIVE inside a paragraph whose
    mark had been deleted, and the op would still report applied.
    """
    case = "delete_block_reaches_runs_nested_in_containers"
    body = (
        "<w:p>"
        "<w:r><w:t>Payment is due within </w:t></w:r>"
        '<w:ins w:id="90" w:author="counterparty" w:date="2026-01-01T00:00:00Z">'
        "<w:r><w:t>forty-five (45)</w:t></w:r></w:ins>"
        "<w:r><w:t> days of </w:t></w:r>"
        '<w:fldSimple w:instr=" DOCPROPERTY Effective ">'
        "<w:r><w:t>the Effective Date</w:t></w:r></w:fldSimple>"
        "<w:r><w:t>.</w:t></w:r>"
        "</w:p>"
    )
    docx_bytes = _build_raw_docx(_raw_heading_p("Section 1. Payment") + body)
    proven, block_map = _prove(
        docx_bytes,
        ops=[{"op": "delete_block", "block_id": "p0001", "issue_key": "PAY-1"}],
    )
    if proven["status"] != "proven":
        failures.append(f"[{case}] fixture transcript did not prove: {proven['failures']!r}")
        return
    for fragment in ("forty-five (45)", "the Effective Date"):
        if fragment not in block_map["p0001"]["text"]:
            failures.append(
                f"[{case}] the fixture's nested run {fragment!r} is not part of the "
                f"block's operative text ({block_map['p0001']['text']!r}) -- the "
                f"fixture does not exercise what it claims to"
            )
            return

    result = redline_block_apply.apply_block_transcript(
        docx_bytes, proven, author=AUTHOR, timestamp_iso=TIMESTAMP
    )
    if result["failures"]:
        failures.append(f"[{case}] unexpected failures: {result['failures']!r}")
    out = result["docx_bytes"]
    if not out:
        failures.append(f"[{case}] no docx_bytes returned")
        return
    _round_trips(case, out, failures)

    target = _paragraphs(out)[1]
    surviving = "".join(t.text or "" for t in target.iter(_qn("t")))
    if surviving:
        failures.append(
            f"[{case}] {surviving!r} SURVIVED the whole-block deletion -- a run nested "
            f"inside <w:ins>/<w:fldSimple> was not wrapped in a <w:del>"
        )
    if _paragraph_mark_revision(target, "del") is None:
        failures.append(f"[{case}] the paragraph mark was not deleted")
    # The counterparty's own pending insertion is not re-authored: our
    # <w:del> nests INSIDE their <w:ins>, which is how Word records
    # "inserted, then deleted before either was accepted".
    counterparty_ins = [
        ins for ins in target.iter(_qn("ins")) if ins.get(_qn("author")) == "counterparty"
    ]
    if len(counterparty_ins) != 1:
        failures.append(f"[{case}] the counterparty's <w:ins> was rewritten or dropped")
    elif not list(counterparty_ins[0].iter(_qn("del"))):
        failures.append(f"[{case}] the counterparty's inserted run was not deleted")


# ---------------------------------------------------------------------------
# AC 3: a block op aimed at a table-cell block fails structured
# ---------------------------------------------------------------------------


def test_block_ops_in_a_table_fail_structured(failures: list) -> None:
    case = "block_ops_in_a_table_fail_structured"
    docx_bytes = _make_table_docx()
    proven, block_map = _prove(
        docx_bytes,
        ops=[
            {"op": "delete_block", "block_id": "p0002", "issue_key": "CELL-1"},
            {
                "op": "insert_block_after",
                "anchor_block_id": "p0002",
                "new_text": _NEW_CLAUSE,
                "issue_key": "CELL-2",
            },
            # A body-paragraph insertion in the SAME transcript: one op's
            # structured failure must never block the others.
            {
                "op": "insert_block_after",
                "anchor_block_id": "p0001",
                "new_text": _NEW_CLAUSE,
                "issue_key": "AUDIT-1",
            },
        ],
    )
    if proven["status"] != "proven":
        failures.append(f"[{case}] fixture transcript did not prove: {proven['failures']!r}")
        return
    if block_map["p0002"]["text"] != _CELL_BODY:
        failures.append(f"[{case}] the fixture's table block is {block_map['p0002']!r}")
        return

    result = redline_block_apply.apply_block_transcript(
        docx_bytes, proven, author=AUTHOR, timestamp_iso=TIMESTAMP
    )

    table_failures = [
        f
        for f in result["failures"]
        if f["reason"] == redline_block_apply.REASON_BLOCK_OP_UNSUPPORTED_IN_TABLE
    ]
    if len(table_failures) != 2:
        failures.append(
            f"[{case}] expected both table-cell block ops to fail structured, got "
            f"{result['failures']!r}"
        )
        return
    for entry in table_failures:
        # The failure names the op AND the block id.
        if entry["block_id"] != "p0002":
            failures.append(f"[{case}] failure does not name the block: {entry!r}")
        if entry["kind"] not in ("delete_block", "insert_block_after"):
            failures.append(f"[{case}] failure does not name the op: {entry!r}")
        if "p0002" not in entry["detail"] or entry["kind"] not in entry["detail"]:
            failures.append(f"[{case}] failure detail names neither op nor block: {entry!r}")
    if {entry["kind"] for entry in table_failures} != {"delete_block", "insert_block_after"}:
        failures.append(f"[{case}] both failures are for the same op: {table_failures!r}")

    # The body-paragraph op in the same transcript still applied ...
    out = result["docx_bytes"]
    if not out:
        failures.append(f"[{case}] one op's table failure blocked the whole batch")
        return
    _round_trips(case, out, failures)
    if [entry["issue_key"] for entry in result["applied"]] != ["AUDIT-1"]:
        failures.append(f"[{case}] applied entries {result['applied']!r}")
    # ... and the table cell itself was left completely alone.
    cell_paragraph = _paragraphs(out)[-1]
    if list(cell_paragraph.iter(_qn("del"))) or list(cell_paragraph.iter(_qn("ins"))):
        failures.append(f"[{case}] the refused op still wrote markup into the table cell")


def test_span_edit_inside_a_table_cell_still_applies(failures: list) -> None:
    case = "span_edit_inside_a_table_cell_still_applies"
    docx_bytes = _make_table_docx()
    proven, _ = _prove(
        docx_bytes,
        patches=[
            {
                "block_id": "p0002",
                "segments": [
                    {"op": "keep", "text": "Deliverables are accepted on "},
                    {"op": "delete", "text": "receipt", "issue_key": "CELL-1"},
                    {"op": "insert", "text": "acceptance testing", "issue_key": "CELL-1"},
                    {"op": "keep", "text": "."},
                ],
            }
        ],
    )
    if proven["status"] != "proven":
        failures.append(f"[{case}] fixture transcript did not prove: {proven['failures']!r}")
        return

    result = redline_block_apply.apply_block_transcript(
        docx_bytes, proven, author=AUTHOR, timestamp_iso=TIMESTAMP
    )
    if result["failures"]:
        failures.append(
            f"[{case}] a WITHIN-CELL span edit was refused -- the table rule must be "
            f"scoped to whole-block ops: {result['failures']!r}"
        )
    out = result["docx_bytes"]
    if not out:
        failures.append(f"[{case}] no docx_bytes returned")
        return
    _round_trips(case, out, failures)

    cell_paragraph = _paragraphs(out)[-1]
    deleted = "".join(t.text or "" for t in cell_paragraph.iter(_qn("delText")))
    inserted = "".join(
        t.text or ""
        for ins in cell_paragraph.iter(_qn("ins"))
        for t in ins.iter(_qn("t"))
    )
    if deleted != "receipt" or inserted != "acceptance testing":
        failures.append(f"[{case}] cell revisions: deleted={deleted!r} inserted={inserted!r}")


# ---------------------------------------------------------------------------
# Block ops and span edits in one transcript
# ---------------------------------------------------------------------------


def test_block_ops_and_span_edits_in_one_transcript(failures: list) -> None:
    case = "block_ops_and_span_edits_in_one_transcript"
    docx_bytes = _make_sectioned_docx(
        [
            ("Section 1. Term", [_TERM_BODY]),
            ("Section 2. Fees", [_FEES_BODY]),
            ("Section 3. Renewal", [_RENEWAL_BODY]),
        ]
    )
    before = _paragraph_texts(docx_bytes)
    proven, _ = _prove(
        docx_bytes,
        patches=[
            {
                "block_id": "p0001",
                "segments": [
                    {"op": "keep", "text": "The Term shall be "},
                    {"op": "delete", "text": "sixty (60)", "issue_key": "TERM-1"},
                    {"op": "insert", "text": "thirty (30)", "issue_key": "TERM-1"},
                    {"op": "keep", "text": " days."},
                ],
            }
        ],
        ops=[
            {
                "op": "insert_block_after",
                "anchor_block_id": "p0002",
                "new_text": _NEW_CLAUSE,
                "issue_key": "AUDIT-1",
            },
            {"op": "delete_block", "block_id": "p0003", "issue_key": "RENEW-1"},
        ],
    )
    if proven["status"] != "proven":
        failures.append(f"[{case}] fixture transcript did not prove: {proven['failures']!r}")
        return

    result = redline_block_apply.apply_block_transcript(
        docx_bytes,
        proven,
        author=AUTHOR,
        timestamp_iso=TIMESTAMP,
        rationale_by_issue={
            "TERM-1": "Term shortened to the negotiated ceiling.",
            "AUDIT-1": "An audit right is required by the negotiating position.",
            "RENEW-1": "Automatic renewal is outside the mandate.",
        },
    )
    if result["failures"]:
        failures.append(f"[{case}] unexpected failures: {result['failures']!r}")
    out = result["docx_bytes"]
    if not out:
        failures.append(f"[{case}] no docx_bytes returned")
        return
    _round_trips(case, out, failures)

    if {entry["kind"] for entry in result["applied"]} != {
        "replace",
        "insert_block_after",
        "delete_block",
    }:
        failures.append(f"[{case}] applied kinds {[e['kind'] for e in result['applied']]!r}")

    # 0 h1, 1 body1 (edited), 2 h2, 3 body2, 4 NEW, 5 h3, 6 body3 (deleted).
    # The insertion renumbers <w:p> 4 onward -- if the deletion had been
    # applied against the post-insertion numbering it would have landed on
    # Section 3's HEADING instead.
    paragraphs = _paragraphs(out)
    if len(paragraphs) != len(before) + 1:
        failures.append(f"[{case}] paragraph count {len(paragraphs)} != {len(before) + 1}")
        return
    inserted_text = "".join(t.text or "" for t in paragraphs[4].iter(_qn("t")))
    if inserted_text != _NEW_CLAUSE:
        failures.append(f"[{case}] <w:p> 4 is {inserted_text!r}, not the inserted clause")
    if not list(paragraphs[6].iter(_qn("delText"))):
        failures.append(f"[{case}] Section 3's BODY was not the paragraph deleted")
    if list(paragraphs[5].iter(_qn("delText"))):
        failures.append(
            f"[{case}] the deletion landed on Section 3's HEADING -- the insertion's "
            f"renumbering corrupted its target"
        )

    with tempfile.TemporaryDirectory() as tmp:
        tmp_path = Path(tmp)
        accepted = [
            text
            for text in _visible_paragraph_texts(_resolved(out, tmp_path, "accept", "accept_all"))
            if text
        ]
        want = [
            "Section 1. Term",
            "The Term shall be thirty (30) days.",
            "Section 2. Fees",
            _FEES_BODY,
            _NEW_CLAUSE,
            "Section 3. Renewal",
        ]
        if accepted != want:
            failures.append(f"[{case}] accept_all text {accepted!r} != {want!r}")
        rejected = [
            text
            for text in _visible_paragraph_texts(_resolved(out, tmp_path, "reject", "reject_all"))
            if text
        ]
        if rejected != [text for text in before if text]:
            failures.append(f"[{case}] reject_all text {rejected!r} != {before!r}")


# ---------------------------------------------------------------------------
# Runner
# ---------------------------------------------------------------------------

TESTS = [
    test_insert_block_after_anchor,
    test_insert_block_at_start,
    test_insert_lands_after_the_anchors_last_physical_paragraph,
    test_two_inserts_after_one_anchor_keep_transcript_order,
    test_delete_block_deletes_runs_and_paragraph_mark,
    test_delete_block_covers_every_physical_paragraph,
    test_delete_block_reaches_runs_nested_in_containers,
    test_block_ops_in_a_table_fail_structured,
    test_span_edit_inside_a_table_cell_still_applies,
    test_block_ops_and_span_edits_in_one_transcript,
]


def main() -> int:
    failures: list = []
    for test in TESTS:
        before = len(failures)
        try:
            test(failures)
        except Exception as exc:  # noqa: BLE001
            failures.append(f"[{test.__name__}] raised {type(exc).__name__}: {exc}")
        if len(failures) == before:
            print(f"PASS: {test.__name__}")
        else:
            for f in failures[before:]:
                print(f"FAIL: {f}")

    print()
    if failures:
        print(f"FAIL: {len(failures)} issue(s) found.")
        return 1
    print("PASS: all redline block-op (issue #622) assertions satisfied.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
