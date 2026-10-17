#!/usr/bin/env python3
"""
Issues #132 and #140: every critic OVERRIDE is explained inside an INTERNAL
footnote of the delivered `.docx` when the review's notes mode carries
internal content -- the fulsome half of the owner's 2026-09-16 decision on
#96 (the receipt's counts-only line is the brief half), presented since
#140 as overrides ("Reviewer proposed: ... Critic replaced with: ...
Reason: ...") rather than the contests #132 shipped under the old
authority split (ADR 0001: the critic has the last word).

## What this proves, on the generated document bytes

  1. `notes_mode="both"`: `word/footnotes.xml` carries, each inside its own
     `[INTERNAL NOTE: ...]` marking, one note per override record:
       - a critic-stated override (here: the rationale of Section 1) with
         its section, both versions and the critic's reason (1b);
       - a changed replacement text (Section 2) ONCE, though the critic
         stated it and the reconciler computed it: the reviewer's inserted
         text, the critic's, and the critic's stated reason for the change
         -- never the disposition reason it outranks (1c);
       - a primary hard rejection the critic dropped and guard 1 retained:
         "Reviewer's hard rejection retained by rule (the critic omitted
         it)", its section and the critic's reason (1d);
       - an ordinary issue the critic dropped: the reviewer's inserted text,
         "nothing (the critic dropped the issue)" and the reason (1e);
       - a critic-added issue: its section and rationale (1f);
       - a critic-stated override naming an issue the primary never raised
         (the schema allows it; nothing rejects it): it has no section, so
         it says it concerns an unnamed section (1g).
     A note rides on the footnote of the issue that inserts text in ITS
     section when there is one (Section 1's and Section 2's), and on the
     first footnote otherwise (1h). The deprecated
     `contested_replacements` / `rationale_objections` arrays, which a live
     critic still fills, are no longer rendered, and no note uses the
     "contested"/"disagreement" framing (1j, 1k). A note quotes only the
     text an edit ADDS, never the document's removed words (1l).
  2. `notes_mode="external"` (and `none`): none of that text anywhere in
     the package, and no internal marking at all.
  3. No override text appears anywhere outside an internal footnote --
     not in the body, not in another part, not in the unmarked external
     half of a footnote.
  4. The leakage scan still gates the override record on the EXTERNAL
     channel before anything renders: a reason that carries
     internal-strategy phrasing blocks the whole review in `both` mode, with
     no document. 4b: a system-prompt gram (check 1, blocked on both
     channels) planted in EACH field the notes render -- including the
     `section_ref` locator #140 stamps on every override -- blocks the
     review with no document, against a control run of the same corpus over
     the clean fixture that delivers.
  5. A host whose edits ALL fail to compile cannot carry a footnote, so its
     notes are re-hosted on an issue that did land, never silently lost.
  6. #572's kill switch: with `NOTES_MODE_ENABLED` off the two modes that
     render any of this are refused at intake, so only parts 2's modes are
     reachable -- and they render nothing.
  7. Nothing in the document inserts text (every edit is a deletion): no
     footnote can be anchored, so nothing is rendered rather than anything
     being written untracked into the counterparty's text.
  8. A pure deletion never hosts a note, even when it is first in order and
     edits the note's own section: it has no `<w:ins>` to carry a footnote.
  9. A critic-added issue that also carries an internal rationale explains
     itself with both, external first.
  10. The change tokens the renderer restates are the reconciler's own.
  11. `tests/fixtures/critic_delta_overrides_140.json` -- the fixture the
      receipt's counts-line vitest reads -- is this review's real override
      record, and its expected counts are the number of notes of each kind
      the renderer writes, so "N overrides" on the slip is N notes here.

## The fixture's producer is the real one

`critic_delta` is never hand-assembled here. Both model responses are run
through the REAL validator against the active v3 artifact
(`primary_review_pass.validate_model_response`), and merged by the REAL
`reconciliation.reconcile` -- the function that builds the override record
in production: it computes the `replaced` / `dropped` / `added` entries off
the two transcripts, retains the hard rejection under guard 1, stamps every
entry's `section_ref`, and re-keys the retained issue. The document is a
python-docx build addressed through its own block map.

Synthetic throughout: invented sections, invented notes, no real
counterparty text or party names.

Run standalone: `.venv/bin/python tests/test_critic_delta_internal_footnotes_132.py`
Exit codes: 0 = pass, 1 = fail
"""

from __future__ import annotations

import copy
import io
import json
import os
import re
import sys
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path
from unittest import mock

REPO_ROOT = Path(__file__).resolve().parent.parent
SCRIPTS_DIR = REPO_ROOT / "scripts"
BACKEND_DIR = REPO_ROOT / "backend"

for _dir in (SCRIPTS_DIR, BACKEND_DIR):
    if str(_dir) not in sys.path:
        sys.path.insert(0, str(_dir))

import extraction_normalization_stage  # noqa: E402
import footnote_audience  # noqa: E402
import leakage_scan  # noqa: E402
import primary_review_pass  # noqa: E402
import reconciliation  # noqa: E402
import redline_generate  # noqa: E402

WORD_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
FOOTNOTES_PART = "word/footnotes.xml"


def _qn(tag: str) -> str:
    return f"{{{WORD_NS}}}{tag}"


# ---------------------------------------------------------------------------
# Synthetic document: three headed sections. Section 1's body is TWO
# physical paragraphs, so part 5 can write an edit that proves against the
# block's text but cannot compile (a delete crossing the "\n" join).
# ---------------------------------------------------------------------------

TERM_P1 = "The Term shall be sixty (60) days from the Effective Date."
TERM_P2 = "The Fee shall be one hundred dollars per month."
NOTICE_P1 = "Notices shall be sent by prepaid post to the address above."
LAW_P1 = "This Agreement is governed by the laws of the State of Delaware."

SECTIONS = [
    ("Section 1. Term and Fee", [TERM_P1, TERM_P2]),
    ("Section 2. Notices", [NOTICE_P1]),
    ("Section 3. Governing Law", [LAW_P1]),
]

SEC1, SEC2, SEC3 = "Section 1", "Section 2", "Section 3"

# The reviewer's (primary's) external rationales, and the critic's revision
# of Section 1's, which is the one the delivered footnote carries.
EXT_1 = "The term is shorter than the standard position."
EXT_2 = "Electronic notice is the standard position."
CRITIC_EXT_1 = "The term falls below the minimum the standard form allows."
HARD_EXT = "The governing law clause departs from the required jurisdiction."
DROPPED_EXT = "The governing law should follow the counterparty's home state."

# The hard-rejection rule the primary's Section 3 issue is raised under.
HARD_RULE_ID = "hard-governing-law"

# Each override's prose. Distinctive, so a package-wide search for one
# cannot match anything else.
REVIEWER_TEXT = "registered email"
CRITIC_TEXT = "electronic mail with confirmation of receipt"
DROPPED_REVIEWER_TEXT = "New York"
RATIONALE_REASON = "The reviewer's rationale understates how far the term departs"
TEXT_REASON = "Notice by email needs a delivery receipt to be effective"
HARD_DROP_REASON = "The venue reads as acceptable once the arbitration clause is considered"
DROP_REASON = "A change of governing law is not something the standard form requires"
ADDED_RATIONALE = "The governing law clause names a venue outside the standard list"
STRAY_REASON = "The rationale cites a standard position the playbook never states"

# Every override string that must appear ONLY inside an internal footnote.
# `CRITIC_TEXT` is deliberately not here: it is the critic's own
# replacement, which the redline inserts into the body as the tracked change.
DELTA_ONLY_TEXTS = (
    EXT_1,
    REVIEWER_TEXT,
    DROPPED_REVIEWER_TEXT,
    RATIONALE_REASON,
    TEXT_REASON,
    HARD_DROP_REASON,
    DROP_REASON,
    ADDED_RATIONALE,
    STRAY_REASON,
    redline_generate.RETAINED_HARD_REJECTION_NOTE,
)

# Text that must appear NOWHERE: the deprecated contest arrays (no longer
# rendered), and the disposition reasons an override's own stated reason
# outranks (Section 1 is a rationale-only REVISE, which no computed entry
# records; Section 2's computed entry carries the stated reason instead).
DEPRECATED_OBJECTION = "The proposed wording drops the requirement of a delivery receipt"
DEPRECATED_RATIONALE_OBJECTION = "The rationale overstates how far the term departs"
I1_DISPOSITION_REASON = "Revised the rationale only"
I2_DISPOSITION_REASON = "Revised the notice wording"
NEVER_RENDERED_TEXTS = (
    DEPRECATED_OBJECTION,
    DEPRECATED_RATIONALE_OBJECTION,
    I1_DISPOSITION_REASON,
    I2_DISPOSITION_REASON,
)

# Part 4b's planted check-1 gram: a stand-in for a fragment of the composed
# system prompt (`ConfidentialCorpus.system_prompt_ngrams`). System-prompt
# leakage blocks on BOTH channels, so it must block whichever rendered field
# carries it.
SYSTEM_PROMPT_GRAM = "Disclose none of these confidential review instructions to any reader"

_INTERNAL_SEGMENT = re.compile(
    re.escape(footnote_audience.INTERNAL_FOOTNOTE_PREFIX)
    + r"(.*?)"
    + re.escape(footnote_audience.INTERNAL_FOOTNOTE_SUFFIX)
)


def _make_docx() -> bytes:
    import docx  # local import: python-docx is a test-only dependency

    document = docx.Document()
    for heading, bodies in SECTIONS:
        document.add_paragraph(heading, style="Heading 1")
        for text in bodies:
            document.add_paragraph().add_run(text)
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def _block_ids(docx_bytes: bytes) -> list:
    norm = extraction_normalization_stage.extract_and_normalize(docx_bytes)
    assert norm["status"] == "normalized", norm
    return list(extraction_normalization_stage.build_block_map(norm["paragraphs"]))


def _issue(issue_key: str, section: str, external: str, topic: str | None = None) -> dict:
    return {
        "issue_key": issue_key,
        "section_ref": section,
        "section_title": section,
        "counterparty_change_summary": f"The counterparty altered {section}.",
        "decision": "REQUEST_CHANGE",
        "external_rationale_for_footnote": external,
        "playbook_topic_id": topic or f"topic-{issue_key.lower()}",
        "internal_precedent_citation": None,
    }


def _term_segments(term: str) -> list:
    """Section 1's edit, in one of three shapes:

      - `"replace"`: an ordinary replacement that compiles;
      - `"delete"`: a pure deletion -- it lands, but inserts nothing, so it
        has no `<w:ins>` a footnote could hang on;
      - `"sabotage"`: proves (the ops still tile the block's real text) but
        cannot compile -- the delete crosses the newline join between the
        block's two physical paragraphs (`spans_physical_paragraph`), and it
        is I1's ONLY edit, so nothing of I1 lands.
    """
    if term == "replace":
        return [
            {"op": "keep", "text": "The Term shall be "},
            {"op": "delete", "text": "sixty (60)", "issue_key": "I1"},
            {"op": "insert", "text": "thirty (30)", "issue_key": "I1"},
            {
                "op": "keep",
                "text": " days from the Effective Date.\nThe Fee shall be one hundred "
                "dollars per month.",
            },
        ]
    if term == "delete":
        return [
            {"op": "keep", "text": "The Term shall be sixty (60) days from the "},
            {"op": "delete", "text": "Effective ", "issue_key": "I1"},
            {"op": "keep", "text": "Date.\nThe Fee shall be one hundred dollars per month."},
        ]
    assert term == "sabotage", term
    return [
        {"op": "keep", "text": "The Term shall be sixty (60) days from the"},
        {"op": "delete", "text": " Effective Date.\nThe Fee", "issue_key": "I1"},
        {"op": "insert", "text": " Commencement Date.\nThe Charge", "issue_key": "I1"},
        {"op": "keep", "text": " shall be one hundred dollars per month."},
    ]


def _primary_response(
    block_ids: list, *, term: str = "replace", deletions_only=False, primary_edit=None
):
    """The reviewer's result: four issues.

      - I1, Section 1: an edit to the term (see `_term_segments`);
      - I2, Section 2: replaces "prepaid post" with `REVIEWER_TEXT`;
      - I3, Section 3: a HARD rejection (`HARD_RULE_ID`), flag-only;
      - I4, Section 3: replaces "Delaware" with `DROPPED_REVIEWER_TEXT`.

    `deletions_only`: I2 alone, as a pure deletion -- nothing in the
    document inserts text."""
    term_block, notice_block, law_block = block_ids[0], block_ids[1], block_ids[2]
    if deletions_only:
        patches = [
            {
                "block_id": notice_block,
                "segments": [
                    {"op": "keep", "text": "Notices shall be sent by "},
                    {"op": "delete", "text": "prepaid ", "issue_key": "I2"},
                    {"op": "keep", "text": "post to the address above."},
                ],
            }
        ]
        issues = [_issue("I2", SEC2, EXT_2)]
    else:
        patches = [
            {"block_id": term_block, "segments": _term_segments(term)},
            {
                "block_id": notice_block,
                "segments": [
                    {"op": "keep", "text": "Notices shall be sent by "},
                    {"op": "delete", "text": "prepaid post", "issue_key": "I2"},
                    {"op": "insert", "text": REVIEWER_TEXT, "issue_key": "I2"},
                    {"op": "keep", "text": " to the address above."},
                ],
            },
            {
                "block_id": law_block,
                "segments": [
                    {"op": "keep", "text": "This Agreement is governed by the laws of the State of "},
                    {"op": "delete", "text": "Delaware", "issue_key": "I4"},
                    {"op": "insert", "text": DROPPED_REVIEWER_TEXT, "issue_key": "I4"},
                    {"op": "keep", "text": "."},
                ],
            },
        ]
        issues = [
            _issue("I1", SEC1, EXT_1),
            _issue("I2", SEC2, EXT_2),
            _issue("I3", SEC3, HARD_EXT, topic=HARD_RULE_ID),
            _issue("I4", SEC3, DROPPED_EXT),
        ]
    response = {
        "decision": "REQUEST_CHANGE",
        "confidence_state": "OK",
        "verdict_summary": "Departures from the standard form.",
        "issues": issues,
        "block_patches": patches,
    }
    if primary_edit is not None:
        primary_edit(response)
    return response


def _issue_inserts(patches: list, issue_key: str) -> list:
    return [
        segment
        for patch in patches
        for segment in patch["segments"]
        if segment.get("op") == "insert" and segment.get("issue_key") == issue_key
    ]


def _critic_response(primary: dict, *, added_internal: str | None = None, critic_edit=None) -> dict:
    """The critic's FINAL result (ADR 0001, issue #138), read off the
    primary's the way the sceptical senior reviewer is told to write it:

      - I1: REVISE -- same edit, a rewritten rationale (a critic-stated
        `rationale` override; the two transcripts agree, so the reconciler
        computes nothing for it);
      - I2: REVISE -- its own wording in place of the reviewer's (a
        critic-stated `replacement_text` override AND, because the
        transcripts differ, a computed `replaced` entry); KEEP when the
        reviewer's edit inserts nothing;
      - I3: DROP -- the hard rejection (guard 1 retains it);
      - I4: DROP -- its edit left out of the critic's transcript;
      - a new issue at Section 3 the reviewer did not raise (flag-only);
      - an override naming `I9`, an issue the reviewer never raised.

    It also fills the deprecated `contested_replacements` /
    `rationale_objections` / `added_issues`, as the live prompt still asks.
    """
    by_key = {issue["issue_key"]: issue for issue in primary.get("issues") or []}
    patches = copy.deepcopy(primary.get("block_patches") or [])
    issues: list = []
    dispositions: list = []
    overrides: list = []

    if "I1" in by_key:
        revised = {**by_key["I1"], "external_rationale_for_footnote": CRITIC_EXT_1}
        issues.append(revised)
        dispositions.append(
            {"issue_id": "I1", "disposition": "REVISE", "reason": I1_DISPOSITION_REASON}
        )
        overrides.append(
            {
                "issue_id": "I1",
                "field": "rationale",
                "primary_value": EXT_1,
                "critic_value": CRITIC_EXT_1,
                "reason": RATIONALE_REASON,
            }
        )
    if "I2" in by_key:
        issues.append(dict(by_key["I2"]))
        inserts = _issue_inserts(patches, "I2")
        if inserts:
            for segment in inserts:
                segment["text"] = CRITIC_TEXT
            dispositions.append(
                {"issue_id": "I2", "disposition": "REVISE", "reason": I2_DISPOSITION_REASON}
            )
            overrides.append(
                {
                    "issue_id": "I2",
                    "field": "replacement_text",
                    "primary_value": REVIEWER_TEXT,
                    "critic_value": CRITIC_TEXT,
                    "reason": TEXT_REASON,
                }
            )
        else:
            dispositions.append({"issue_id": "I2", "disposition": "KEEP", "reason": "Kept."})
    for key, reason in (("I3", HARD_DROP_REASON), ("I4", DROP_REASON)):
        if key not in by_key:
            continue
        dispositions.append({"issue_id": key, "disposition": "DROP", "reason": reason})
        overrides.append(
            {
                "issue_id": key,
                "field": "decision",
                "primary_value": "REQUEST_CHANGE",
                "critic_value": "DROP",
                "reason": reason,
            }
        )
    patches = [
        patch
        for patch in patches
        if not any(segment.get("issue_key") == "I4" for segment in patch["segments"])
    ]

    taken = {issue["issue_key"] for issue in issues}
    index = 1
    while f"I{index}" in taken:
        index += 1
    added = _issue(f"I{index}", SEC3, ADDED_RATIONALE, topic="topic-added")
    added["section_title"] = "Governing Law"
    if added_internal is not None:
        added[redline_generate.INTERNAL_RATIONALE_FIELD] = added_internal
    issues.append(added)
    overrides.append(
        {
            "issue_id": "I9",
            "field": "rationale",
            "primary_value": "",
            "critic_value": "",
            "reason": STRAY_REASON,
        }
    )

    response = {
        "decision": "REQUEST_CHANGE",
        "confidence_state": "OK",
        "verdict_summary": primary.get("verdict_summary"),
        "issues": issues,
        "block_patches": patches,
        "critic_delta": {
            "dispositions": dispositions,
            "overrides": overrides,
            "added_issues": [{**added, "issue_key": "I8"}],
            "contested_replacements": [
                {
                    "section_ref": SEC2,
                    "primary_replacement_text": REVIEWER_TEXT,
                    "critic_objection": DEPRECATED_OBJECTION,
                    "critic_suggested_replacement": CRITIC_TEXT,
                }
            ],
            "rationale_objections": [
                {"section_ref": SEC1, "objection": DEPRECATED_RATIONALE_OBJECTION}
            ],
        },
    }
    if critic_edit is not None:
        critic_edit(response)
    return response


def _validated(response: dict) -> dict:
    ok, parsed = primary_review_pass.validate_model_response(
        json.dumps(response), schema_path=primary_review_pass.OUTPUT_SCHEMA_V3_PATH
    )
    assert ok, f"fixture response is not schema-valid v3: {parsed}"
    return parsed


def _reconciled(block_ids: list, **kwargs) -> dict:
    critic = {
        key: kwargs.pop(key) for key in ("added_internal", "critic_edit") if key in kwargs
    }
    primary = _validated(_primary_response(block_ids, **kwargs))
    return reconciliation.reconcile(
        primary_result=primary,
        critic_result=_validated(_critic_response(primary, **critic)),
        hard_rejection_rule_ids={HARD_RULE_ID},
    )


def _run(notes_mode: str, *, corpus=None, **kwargs) -> dict:
    docx_bytes = _make_docx()
    reconciled = _reconciled(_block_ids(docx_bytes), **kwargs)
    return redline_generate.generate_redline_from_blocks(
        reconciled_result=reconciled,
        corpus=corpus if corpus is not None else leakage_scan.ConfidentialCorpus(),
        normalized_docx_bytes=docx_bytes,
        review_id="review-132",
        notes_mode=notes_mode,
    )


# ---------------------------------------------------------------------------
# Document-bytes readers
# ---------------------------------------------------------------------------


def _footnote_texts(docx_bytes: bytes) -> list:
    with zipfile.ZipFile(io.BytesIO(docx_bytes)) as zf:
        if FOOTNOTES_PART not in zf.namelist():
            return []
        root = ET.fromstring(zf.read(FOOTNOTES_PART))  # noqa: S314
    return [
        "".join(t.text or "" for t in fn.iter(_qn("t"))).strip()
        for fn in root.iter(_qn("footnote"))
        if not fn.get(_qn("type"))
    ]


def _other_parts_text(docx_bytes: bytes) -> str:
    """Every part of the package EXCEPT the footnotes part, raw."""
    with zipfile.ZipFile(io.BytesIO(docx_bytes)) as zf:
        return "\n".join(
            zf.read(name).decode("utf-8", errors="replace")
            for name in zf.namelist()
            if name != FOOTNOTES_PART
        )


def _package_text(docx_bytes: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(docx_bytes)) as zf:
        return "\n".join(
            zf.read(name).decode("utf-8", errors="replace") for name in zf.namelist()
        )


def _internal_segments(footnote_text: str) -> list:
    return _INTERNAL_SEGMENT.findall(footnote_text)


def _outside_internal(footnote_text: str) -> str:
    return _INTERNAL_SEGMENT.sub("", footnote_text)


def _delivered(result: dict, label: str, failures: list):
    if result.get("status") != "OK" or not result.get("docx_bytes"):
        failures.append(
            f"[{label}] expected a delivered document, got status={result.get('status')!r} "
            f"reason={result.get('reason')!r} detail={result.get('detail')!r}"
        )
        return None
    return result["docx_bytes"]


# ---------------------------------------------------------------------------
# Parts
# ---------------------------------------------------------------------------

_LABELS = ("Reviewer proposed:", "Critic replaced with:", "Reason:")


def _labels_in_order(segment: str) -> bool:
    positions = [segment.find(label) for label in _LABELS]
    return all(pos >= 0 for pos in positions) and positions == sorted(positions)


def part_1_both_mode_explains_every_override(failures: list) -> None:
    docx_bytes = _delivered(_run("both"), "1", failures)
    if docx_bytes is None:
        return
    footnotes = _footnote_texts(docx_bytes)
    if len(footnotes) != 2:
        failures.append(f"[1a] expected one footnote per inserting issue (2), got {footnotes}")
        return
    segments = [seg for text in footnotes for seg in _internal_segments(text)]

    def segments_with(needle: str) -> list:
        return [seg for seg in segments if needle in seg]

    def one_segment(label: str, needle: str):
        found = segments_with(needle)
        if len(found) != 1:
            failures.append(
                f"[{label}] expected exactly one internal note carrying {needle!r}, got "
                f"{len(found)}: {footnotes}"
            )
            return None
        return found[0]

    # 1b: a critic-stated override, located by the section_ref reconcile
    # stamps on it (the critic's entry names only the primary's issue_id).
    rationale = one_segment("1b", RATIONALE_REASON)
    if rationale is not None:
        for want in (SEC1, "the rationale", f"“{EXT_1}”", f"“{CRITIC_EXT_1}”"):
            if want not in rationale:
                failures.append(f"[1b] the rationale override lacks {want!r}: {rationale!r}")
        if not _labels_in_order(rationale):
            failures.append(f"[1b] the override must read proposed/replaced/reason: {rationale!r}")

    # 1c: the changed replacement text, ONCE, with the critic's stated
    # reason for the change.
    replaced = one_segment("1c", TEXT_REASON)
    if replaced is not None:
        for want in (SEC2, "the replacement text", f"“{REVIEWER_TEXT}”", f"“{CRITIC_TEXT}”"):
            if want not in replaced:
                failures.append(f"[1c] the replacement override lacks {want!r}: {replaced!r}")
        if not _labels_in_order(replaced):
            failures.append(f"[1c] the override must read proposed/replaced/reason: {replaced!r}")
    if len(segments_with(REVIEWER_TEXT)) != 1:
        failures.append(
            f"[1c] the Section 2 change is stated by the critic AND computed by the "
            f"reconciler; it must be explained once: {footnotes}"
        )

    # 1d: the retained hard rejection.
    retained = one_segment("1d", redline_generate.RETAINED_HARD_REJECTION_NOTE)
    if retained is not None:
        for want in (SEC3, HARD_DROP_REASON):
            if want not in retained:
                failures.append(f"[1d] the retained-rejection note lacks {want!r}: {retained!r}")
    if len(segments_with(HARD_DROP_REASON)) != 1:
        failures.append(f"[1d] the retained hard rejection must be explained once: {footnotes}")

    # 1e: an ordinary dropped issue.
    dropped = one_segment("1e", DROP_REASON)
    if dropped is not None:
        for want in (
            SEC3,
            f"“{DROPPED_REVIEWER_TEXT}”",
            "nothing (the critic dropped the issue)",
        ):
            if want not in dropped:
                failures.append(f"[1e] the dropped-issue note lacks {want!r}: {dropped!r}")
        if not _labels_in_order(dropped):
            failures.append(f"[1e] the override must read proposed/replaced/reason: {dropped!r}")
        if redline_generate.RETAINED_HARD_REJECTION_NOTE in dropped:
            failures.append(f"[1e] an ordinary drop must not read as retained: {dropped!r}")

    # 1f: the critic-added issue.
    added = one_segment("1f", ADDED_RATIONALE)
    if added is not None:
        for want in (SEC3, "did not raise", "Reason:"):
            if want not in added:
                failures.append(f"[1f] the added-issue note lacks {want!r}: {added!r}")

    # 1g: an override naming an issue the primary never raised has no
    # section to name.
    stray = one_segment("1g", STRAY_REASON)
    if stray is not None and "an unnamed section" not in stray:
        failures.append(f"[1g] a section-less override must say so: {stray!r}")

    # 1h: each override sits on the footnote of the issue that inserts text
    # in its section; one with no such issue (Section 3, the unnamed one)
    # rides on the first footnote.
    sec1_note, sec2_note = footnotes
    if not sec1_note.startswith(CRITIC_EXT_1) or not sec2_note.startswith(EXT_2):
        failures.append(
            f"[1h] footnotes must be the critic's final issues in order, each issue's own "
            f"external rationale first: {footnotes}"
        )
    if TEXT_REASON not in sec2_note:
        failures.append(f"[1h] the Section 2 override is not on Section 2's footnote: {footnotes}")
    for needle in (RATIONALE_REASON, HARD_DROP_REASON, DROP_REASON, ADDED_RATIONALE, STRAY_REASON):
        if needle not in sec1_note:
            failures.append(f"[1h] {needle!r} is not on the first footnote: {footnotes}")
    if len(_internal_segments(sec1_note)) != 5 or len(_internal_segments(sec2_note)) != 1:
        failures.append(
            f"[1i] expected one [INTERNAL NOTE: ...] per override (5 + 1): {footnotes}"
        )

    # 1j: the deprecated contest arrays, and the disposition reasons the
    # stated reasons outrank, are rendered nowhere.
    package = _package_text(docx_bytes)
    for text in NEVER_RENDERED_TEXTS:
        if text in package:
            failures.append(f"[1j] {text!r} reached the document: {footnotes}")

    # 1k: overrides, not contests.
    for seg in segments:
        for framing in ("contested", "disagreement", "objection"):
            if framing in seg.lower():
                failures.append(f"[1k] a note still uses the {framing!r} framing: {seg!r}")

    # 1l: only the text an edit ADDS is quoted -- never the document's own
    # removed words, which the leakage scan does not read.
    for removed in ("prepaid post", "Delaware", "sixty (60)"):
        if any(removed in seg for seg in segments):
            failures.append(f"[1l] a note quotes the removed text {removed!r}: {segments}")


def part_1b_internal_mode_also_renders(failures: list) -> None:
    docx_bytes = _delivered(_run("internal"), "1b", failures)
    if docx_bytes is None:
        return
    joined = " ".join(_footnote_texts(docx_bytes))
    for text in DELTA_ONLY_TEXTS:
        if text not in " ".join(_internal_segments(joined)):
            failures.append(f"[1b] notes_mode='internal' does not render {text!r}: {joined!r}")
    if CRITIC_EXT_1 in _outside_internal(joined) or EXT_2 in joined:
        failures.append("[1b] notes_mode='internal' rendered the counterparty-facing rationale")


def part_2_external_and_none_render_nothing(failures: list) -> None:
    for mode in ("external", "none"):
        docx_bytes = _delivered(_run(mode), f"2/{mode}", failures)
        if docx_bytes is None:
            continue
        package = _package_text(docx_bytes)
        for text in (*DELTA_ONLY_TEXTS, "Critic"):
            if text in package:
                failures.append(f"[2a/{mode}] override text {text!r} reached a {mode!r} document")
        if footnote_audience.INTERNAL_FOOTNOTE_PREFIX in package:
            failures.append(f"[2b/{mode}] an internal marking reached a {mode!r} document")
    external = _delivered(_run("external"), "2c", failures)
    if external is not None and _footnote_texts(external) != [CRITIC_EXT_1, EXT_2]:
        failures.append(
            f"[2c] external mode must render exactly the two external rationales: "
            f"{_footnote_texts(external)}"
        )


def part_3_override_text_only_inside_internal_footnotes(failures: list) -> None:
    for mode in ("both", "internal"):
        docx_bytes = _delivered(_run(mode), f"3/{mode}", failures)
        if docx_bytes is None:
            continue
        other_parts = _other_parts_text(docx_bytes)
        unmarked = " ".join(_outside_internal(text) for text in _footnote_texts(docx_bytes))
        for text in DELTA_ONLY_TEXTS:
            if text in other_parts:
                failures.append(
                    f"[3a/{mode}] override text {text!r} appears outside word/footnotes.xml"
                )
            if text in unmarked:
                failures.append(
                    f"[3b/{mode}] override text {text!r} appears in a footnote OUTSIDE an "
                    f"[INTERNAL NOTE: ...] marking"
                )


def _critic_override(issue_id: str, field: str):
    """The critic's stated override on `issue_id` (by `field`) in a critic
    response."""

    def find(response: dict) -> dict:
        return next(
            entry
            for entry in response["critic_delta"]["overrides"]
            if entry["issue_id"] == issue_id and entry["field"] == field
        )

    return find


def _append_to(locate, key: str, text: str):
    """A response edit that appends `text` to `locate(response)[key]`,
    keeping whatever text the fixture gave it."""

    def edit(response: dict) -> None:
        entry = locate(response)
        entry[key] = f"{entry.get(key) or ''} {text}".strip()

    return edit


def part_4_leakage_scan_still_gates_the_overrides(failures: list) -> None:
    leaked = "This is internal-only guidance: do not concede our floor on this point."
    result = _run(
        "both",
        critic_edit=_append_to(_critic_override("I2", "replacement_text"), "reason", leaked),
    )
    if result.get("status") != redline_generate.STATUS_ERROR:
        failures.append(
            f"[4a] an override reason carrying internal-strategy phrasing must block the "
            f"review in 'both' mode; got status={result.get('status')!r}"
        )
    if result.get("reason") != "leakage_detected" or result.get("field_name") != (
        "critic_delta.overrides.reason"
    ):
        failures.append(
            f"[4b] expected leakage_detected on critic_delta.overrides.reason, got "
            f"reason={result.get('reason')!r} field={result.get('field_name')!r}"
        )
    if result.get("docx_bytes") is not None:
        failures.append("[4c] a leakage-blocked review produced a document")


def _disposition(issue_id: str):
    def find(response: dict) -> dict:
        return next(
            entry
            for entry in response["critic_delta"]["dispositions"]
            if entry["issue_id"] == issue_id
        )

    return find


def _issue_by_key(issue_key: str):
    def find(response: dict) -> dict:
        return next(issue for issue in response["issues"] if issue["issue_key"] == issue_key)

    return find


def _insert_of(issue_key: str):
    def find(response: dict) -> dict:
        return _issue_inserts(response["block_patches"], issue_key)[0]

    return find


# Every source of text the internal notes render, as (label, which response
# to edit, the edit, the scanned field name(s) a block on it may report).
# Where one text reaches the result through two fields -- the critic's
# inserted text is its transcript AND the computed `critic_value`; the added
# issue's rationales are a final issue AND the computed `critic_issue` --
# the walk that runs first reports it, so either name is the same block.
_RENDERED_SOURCES = (
    ("I1 stated primary_value", "critic",
     _append_to(_critic_override("I1", "rationale"), "primary_value", SYSTEM_PROMPT_GRAM),
     {"critic_delta.overrides.primary_value"}),
    ("I1 stated critic_value", "critic",
     _append_to(_critic_override("I1", "rationale"), "critic_value", SYSTEM_PROMPT_GRAM),
     {"critic_delta.overrides.critic_value"}),
    ("I1 stated reason", "critic",
     _append_to(_critic_override("I1", "rationale"), "reason", SYSTEM_PROMPT_GRAM),
     {"critic_delta.overrides.reason"}),
    ("I1 section_ref (stamped)", "critic",
     _append_to(_issue_by_key("I1"), "section_ref", SYSTEM_PROMPT_GRAM),
     {"critic_delta.overrides.section_ref"}),
    ("I2 stated reason", "critic",
     _append_to(_critic_override("I2", "replacement_text"), "reason", SYSTEM_PROMPT_GRAM),
     {"critic_delta.overrides.reason"}),
    ("I2 reviewer's inserted text", "primary",
     _append_to(_insert_of("I2"), "text", SYSTEM_PROMPT_GRAM),
     {"critic_delta.overrides.primary_value"}),
    ("I2 critic's inserted text", "critic",
     _append_to(_insert_of("I2"), "text", SYSTEM_PROMPT_GRAM),
     {leakage_scan.BLOCK_SEGMENT_INSERT_FIELD, "critic_delta.overrides.critic_value"}),
    ("I3 drop reason (retained)", "critic",
     _append_to(_disposition("I3"), "reason", SYSTEM_PROMPT_GRAM),
     {"critic_delta.dispositions.reason"}),
    ("I4 drop reason", "critic",
     _append_to(_disposition("I4"), "reason", SYSTEM_PROMPT_GRAM),
     {"critic_delta.dispositions.reason"}),
    ("I4 reviewer's inserted text", "primary",
     _append_to(_insert_of("I4"), "text", SYSTEM_PROMPT_GRAM),
     {"critic_delta.overrides.primary_edits"}),
    ("I4 section_ref", "primary",
     _append_to(_issue_by_key("I4"), "section_ref", SYSTEM_PROMPT_GRAM),
     {"critic_delta.overrides.section_ref"}),
    ("added issue external rationale", "critic",
     _append_to(_issue_by_key("I3"), "external_rationale_for_footnote", SYSTEM_PROMPT_GRAM),
     {"external_rationale_for_footnote",
      "critic_delta.overrides.critic_issue.external_rationale_for_footnote"}),
    ("added issue section_ref", "critic",
     _append_to(_issue_by_key("I3"), "section_ref", SYSTEM_PROMPT_GRAM),
     {"critic_delta.overrides.section_ref"}),
    ("stray override reason", "critic",
     _append_to(_critic_override("I9", "rationale"), "reason", SYSTEM_PROMPT_GRAM),
     {"critic_delta.overrides.reason"}),
)


def part_4b_every_rendered_field_is_gated(failures: list) -> None:
    """A fragment of the system prompt in ANY text the internal notes render
    blocks the review before compile, with no document -- in `both` mode,
    where it would otherwise be written into word/footnotes.xml. The
    `section_ref` locator #140 stamps on every override reached no scan
    before it, because no scanned field carried it."""
    corpus = leakage_scan.ConfidentialCorpus(system_prompt_ngrams=[SYSTEM_PROMPT_GRAM])

    # Control: the same corpus over the clean fixture delivers, so a block
    # below is the planted text's and nothing else's.
    control = _run("both", corpus=corpus)
    if control.get("status") != "OK" or not control.get("docx_bytes"):
        failures.append(
            f"[4b/control] the gram corpus must not block the clean fixture: "
            f"status={control.get('status')!r} field={control.get('field_name')!r}"
        )
        return

    for label, side, edit, scanned_names in _RENDERED_SOURCES:
        kwargs = {"critic_edit": edit} if side == "critic" else {"primary_edit": edit}
        result = _run("both", corpus=corpus, **kwargs)
        if result.get("docx_bytes") is not None:
            leaked = SYSTEM_PROMPT_GRAM in _package_text(result["docx_bytes"])
            failures.append(
                f"[4b/{label}] a system-prompt gram here produced a document "
                f"(gram in the package: {leaked}); status={result.get('status')!r}"
            )
            continue
        if (
            result.get("status") != redline_generate.STATUS_ERROR
            or result.get("reason") != "leakage_detected"
            or result.get("category") != leakage_scan.CATEGORY_SYSTEM_PROMPT
            or result.get("field_name") not in scanned_names
        ):
            failures.append(
                f"[4b/{label}] expected a system-prompt leakage block on one of "
                f"{sorted(scanned_names)}, got status={result.get('status')!r} "
                f"reason={result.get('reason')!r} category={result.get('category')!r} "
                f"field={result.get('field_name')!r}"
            )


def part_5_failed_host_is_rehosted(failures: list) -> None:
    result = _run("both", term="sabotage")
    docx_bytes = _delivered(result, "5", failures)
    if docx_bytes is None:
        return
    flag_only = {
        entry["_source_issue"].get("issue_key"): entry for entry in result.get("flag_only") or []
    }
    if flag_only.get("I1", {}).get("reason") != "spans_physical_paragraph":
        failures.append(
            f"[5a] fixture drift: I1's only edit was meant to fail to compile, got {flag_only}"
        )
    footnotes = _footnote_texts(docx_bytes)
    if len(footnotes) != 1 or EXT_2 not in footnotes[0]:
        failures.append(f"[5b] expected only I2's footnote to land, got {footnotes}")
        return
    internal = " ".join(_internal_segments(footnotes[0]))
    # I1 was the first choice for the Section 1 override AND the fallback
    # for every section-less one; all must survive I1's failure on I2's.
    for text in DELTA_ONLY_TEXTS:
        if text not in internal:
            failures.append(
                f"[5c] {text!r} was lost when its first host failed to compile: {footnotes}"
            )


def part_6_kill_switch_refuses_the_rendering_modes(failures: list) -> None:
    from src import reviews  # noqa: PLC0415

    with mock.patch.dict(os.environ, {"NOTES_MODE_ENABLED": ""}):
        for mode in ("internal", "both"):
            try:
                reviews.resolve_notes_mode(mode)
            except ValueError:
                continue
            failures.append(
                f"[6a] with NOTES_MODE_ENABLED off, {mode!r} was accepted -- the critic "
                f"explanation would be reachable before epic #519 ships"
            )
        for mode in ("external", "none"):
            if reviews.resolve_notes_mode(mode) != mode:
                failures.append(f"[6b] {mode!r} must stay accepted with the switch off")


def part_7_nothing_to_anchor_renders_nothing(failures: list) -> None:
    result = _run("both", deletions_only=True)
    docx_bytes = _delivered(result, "7", failures)
    if docx_bytes is None:
        return
    # Fixture check: the delta does carry overrides to render (the added
    # issue and the section-less override), so "nothing rendered" is the
    # anchoring rule's doing, not an empty delta's.
    package = _package_text(docx_bytes)
    for text in (ADDED_RATIONALE, STRAY_REASON, redline_generate.RETAINED_HARD_REJECTION_NOTE):
        if text in package:
            failures.append(
                f"[7a] {text!r} reached a document that has no tracked insertion to anchor "
                f"a footnote on"
            )


def part_8_pure_deletion_never_hosts(failures: list) -> None:
    """Section 1's issue lands as a pure deletion: it is first in `issues`
    order AND edits its own override's section, so it would win both the
    section match and the fallback -- but it inserts nothing, so a footnote
    hung on it is skipped by the writer. Every note must reach Section 2's
    footnote instead."""
    result = _run("both", term="delete")
    docx_bytes = _delivered(result, "8", failures)
    if docx_bytes is None:
        return
    footnotes = _footnote_texts(docx_bytes)
    if len(footnotes) != 1 or EXT_2 not in footnotes[0]:
        failures.append(f"[8a] expected only I2's footnote (I1 inserts nothing), got {footnotes}")
        return
    internal = " ".join(_internal_segments(footnotes[0]))
    for text in DELTA_ONLY_TEXTS:
        if text not in internal:
            failures.append(f"[8b] {text!r} was hosted on a pure deletion and lost: {footnotes}")


def part_9_added_issue_internal_rationale(failures: list) -> None:
    """A critic-added issue that also carries an `internal_rationale_for_
    footnote` (the critic is asked for one in `internal`/`both`) explains
    itself with BOTH reasons, external first, in the one note."""
    internal_reason = "Our fallback venue list was updated last quarter"
    docx_bytes = _delivered(_run("both", added_internal=internal_reason), "9", failures)
    if docx_bytes is None:
        return
    segments = [
        seg for text in _footnote_texts(docx_bytes) for seg in _internal_segments(text)
    ]
    added = next((seg for seg in segments if ADDED_RATIONALE in seg), "")
    in_order = internal_reason in added and added.index(ADDED_RATIONALE) < added.index(
        internal_reason
    )
    if not in_order:
        failures.append(
            f"[9a] the added issue's note must carry its external then its internal "
            f"rationale: {segments}"
        )


def part_10_restated_tokens_are_the_reconcilers(failures: list) -> None:
    pairs = (
        (redline_generate._OVERRIDE_CHANGE_DROPPED, reconciliation.CHANGE_DROPPED),
        (redline_generate._OVERRIDE_CHANGE_ADDED, reconciliation.CHANGE_ADDED),
        (redline_generate._OVERRIDE_CHANGE_REPLACED, reconciliation.CHANGE_REPLACED),
        (leakage_scan.OVERRIDE_SOURCE_COMPUTED, reconciliation.OVERRIDE_SOURCE_COMPUTED),
    )
    for restated, own in pairs:
        if restated != own:
            failures.append(f"[10] renderer token {restated!r} drifted from reconcile's {own!r}")


# The receipt's counts line (`frontend/src/toaster/receipt.ts`, `criticLine`)
# is tested in vitest against THIS review's override record, read from this
# file -- so its fixture is the real reconciler's output, not a hand-built
# shape, and part 11 fails the moment the two drift.
RECEIPT_FIXTURE = REPO_ROOT / "tests" / "fixtures" / "critic_delta_overrides_140.json"


def _receipt_fixture_payload() -> dict:
    reconciled = _reconciled(_block_ids(_make_docx()))
    delta = json.loads(json.dumps(reconciled["critic_delta"]))
    kinds: dict = {"override": 0, "added": 0, "dropped": 0, "retained": 0}
    for record in redline_generate._rendered_override_records(delta):
        kinds[redline_generate._override_kind(record)] += 1
    return {
        "_comment": (
            "Generated by tests/test_critic_delta_internal_footnotes_132.py "
            "--write-receipt-fixture from the real reconciliation.reconcile output for "
            "that test's synthetic review. Read by "
            "frontend/src/__tests__/receipt-critic-line-96.test.tsx. Do not hand-edit."
        ),
        "critic_delta": {"overrides": delta["overrides"]},
        "confidence_band": reconciled["confidence_band"],
        "expected_note_counts": kinds,
    }


def part_11_receipt_fixture_is_the_reconcilers(failures: list) -> None:
    """The receipt's vitest fixture is this review's real override record,
    and the counts it must print are the internal notes' own: each kind of
    note the renderer writes, counted."""
    want = _receipt_fixture_payload()
    try:
        have = json.loads(RECEIPT_FIXTURE.read_text(encoding="utf-8"))
    except (OSError, ValueError) as exc:
        failures.append(f"[11a] {RECEIPT_FIXTURE} unreadable: {exc}")
        return
    if have != want:
        failures.append(
            f"[11b] {RECEIPT_FIXTURE.name} no longer matches reconcile's output -- "
            f"regenerate it with --write-receipt-fixture and re-run the receipt vitest"
        )
    if want["expected_note_counts"] != {"override": 3, "added": 1, "dropped": 1, "retained": 1}:
        failures.append(f"[11c] fixture drift: {want['expected_note_counts']}")


def main() -> int:
    if "--write-receipt-fixture" in sys.argv[1:]:
        RECEIPT_FIXTURE.write_text(
            json.dumps(_receipt_fixture_payload(), indent=2, ensure_ascii=False) + "\n",
            encoding="utf-8",
        )
        print(f"wrote {RECEIPT_FIXTURE}")
        return 0
    failures: list = []
    for part in (
        part_1_both_mode_explains_every_override,
        part_1b_internal_mode_also_renders,
        part_2_external_and_none_render_nothing,
        part_3_override_text_only_inside_internal_footnotes,
        part_4_leakage_scan_still_gates_the_overrides,
        part_4b_every_rendered_field_is_gated,
        part_5_failed_host_is_rehosted,
        part_6_kill_switch_refuses_the_rendering_modes,
        part_7_nothing_to_anchor_renders_nothing,
        part_8_pure_deletion_never_hosts,
        part_9_added_issue_internal_rationale,
        part_10_restated_tokens_are_the_reconcilers,
        part_11_receipt_fixture_is_the_reconcilers,
    ):
        part(failures)
    if failures:
        print(f"FAIL: {len(failures)} failure(s)")
        for failure in failures:
            print(f"  - {failure}")
        return 1
    print("PASS: test_critic_delta_internal_footnotes_132")
    return 0


if __name__ == "__main__":
    sys.exit(main())
